"""Экспорт документов: Markdown, HTML, PDF, DOCX из одного Markdown-текста материала.

Текст разбирается один раз в простую модель (заголовки, абзацы, списки, таблицы, цитаты, код), и четыре формата
строятся из неё одинаково. HTML самодостаточный: стили внутри, внешних ресурсов нет. DOCX использует настоящие
стили Word (заголовки доступны для оглавления, таблицы и нумерованные списки — не картинки и не набранные вручную
номера). PDF рисуется собственным шрифтом DejaVu Sans (кириллица не зависит от системных шрифтов), таблицы переносят
текст в ячейках и повторяют шапку на каждой странице.
"""
from __future__ import annotations

import html
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

FORMATS = {
    "md": ("text/markdown; charset=utf-8", "Markdown"),
    "html": ("text/html; charset=utf-8", "HTML"),
    "pdf": ("application/pdf", "PDF"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "DOCX"),
}
FONT_DIR = Path(__file__).with_name("fonts")
MAX_CHARS = 500_000


class ExportError(ValueError):
    """Экспорт не удался; текст понятен пользователю."""


# --------------------------------------------------------------------------- модель документа

@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    href: str = ""


@dataclass
class Block:
    kind: str                       # heading, para, ul, ol, table, quote, code, hr
    runs: list[Run] = field(default_factory=list)          # heading, para
    level: int = 0                                          # heading
    items: list[list[Run]] = field(default_factory=list)   # ul, ol
    header: list[list[Run]] = field(default_factory=list)  # table
    rows: list[list[list[Run]]] = field(default_factory=list)
    lines: list[list[Run]] = field(default_factory=list)   # quote: абзацы
    text: str = ""                                          # code
    title: str = ""                                         # quote (callout)


_INLINE = re.compile(
    r"(?P<code>`[^`\n]+`)|(?P<bold>\*\*(?!\s)(?:.+?)\*\*)|(?P<wiki>\[\[[^\]\n]+\]\])|(?P<link>\[[^\]\n]+\]\([^)\s]+\))"
    r"|(?P<em>(?<![\w*])\*(?!\s)[^*\n]+?\*(?!\w)|(?<![\w_])_(?!\s)[^_\n]+?_(?![\w_]))")


def inline(text: str, bold: bool = False, italic: bool = False) -> list[Run]:
    runs: list[Run] = []
    pos = 0
    for match in _INLINE.finditer(text):
        if match.start() > pos:
            runs.append(Run(text[pos:match.start()], bold, italic))
        token = match.group(0)
        kind = match.lastgroup
        if kind == "code":
            runs.append(Run(token[1:-1], bold, italic, code=True))
        elif kind == "bold":
            runs += inline(token[2:-2], True, italic)
        elif kind == "em":
            runs += inline(token[1:-1], bold, True)
        elif kind == "wiki":
            target, _, alias = token[2:-2].partition("|")
            runs.append(Run(alias or target.rsplit("/", 1)[-1].removesuffix(".md"), bold, italic))
        elif kind == "link":
            label, _, url = token[1:-1].partition("](")
            runs.append(Run(label, bold, italic, href=url if re.match(r"(?i)(https?|mailto|obsidian):", url) else ""))
        pos = match.end()
    if pos < len(text):
        runs.append(Run(text[pos:], bold, italic))
    return [r for r in runs if r.text]


def plain(runs: list[Run]) -> str:
    return "".join(r.text for r in runs)


def _cells(line: str) -> list[str]:
    body = line.strip()
    body = body[1:] if body.startswith("|") else body
    body = body[:-1] if body.endswith("|") and not body.endswith("\\|") else body
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", body)]


def _is_rule(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*", line)) and "-" in line


def parse(markdown: str) -> list[Block]:
    """Markdown в блоки. Понимает то, что пишут материалы Memory Desk и заметки Obsidian: заголовки, абзацы, списки,
    таблицы, цитаты и `> [!callout]`, блоки кода, `---`. Фронтматтер YAML в начале отбрасывается."""
    text = markdown.replace("\r\n", "\n").replace("\r", "\n")
    if text.startswith("---\n"):
        end = text.find("\n---\n", 4)
        if end != -1:
            text = text[end + 5:]
    lines = text.split("\n")
    blocks: list[Block] = []
    para: list[str] = []
    i = 0

    def flush() -> None:
        if para:
            runs: list[Run] = []
            for index, text_line in enumerate(para):
                if index:
                    runs.append(Run("\n"))
                runs += inline(text_line.strip())
            blocks.append(Block("para", runs=runs))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        fence = re.match(r"^(```|~~~)\s*([\w-]*)", stripped)
        if fence:
            flush()
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(fence.group(1)):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append(Block("code", text="\n".join(body), title=fence.group(2)))
            continue
        if not stripped:
            flush()
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", stripped)
        if heading:
            flush()
            blocks.append(Block("heading", level=len(heading.group(1)), runs=inline(heading.group(2))))
            i += 1
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            flush()
            blocks.append(Block("hr"))
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and _is_rule(lines[i + 1]):
            flush()
            header = [inline(c) for c in _cells(stripped)]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [inline(c) for c in _cells(lines[i])]
                cells += [[] for _ in range(len(header) - len(cells))]
                rows.append(cells[:len(header)])
                i += 1
            blocks.append(Block("table", header=header, rows=rows))
            continue
        if stripped.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            title = ""
            callout = re.match(r"\[!(\w+)\][+-]?\s*(.*)", quote[0]) if quote else None
            if callout:
                title = callout.group(2).strip() or callout.group(1).capitalize()
                quote = quote[1:]
            paragraphs, current = [], []
            for q in quote + [""]:
                if q.strip():
                    current.append(q.strip())
                elif current:
                    paragraphs.append(inline(" ".join(current)))
                    current = []
            blocks.append(Block("quote", lines=paragraphs, title=title))
            continue
        bullet = re.match(r"^(\s*)[-*+•]\s+(.*)", line)
        number = re.match(r"^(\s*)\d+[.)]\s+(.*)", line)
        if bullet or number:
            flush()
            ordered = bool(number)
            items: list[list[Run]] = []
            while i < len(lines):
                m = re.match(r"^(\s*)\d+[.)]\s+(.*)", lines[i]) if ordered else re.match(r"^(\s*)[-*+•]\s+(.*)", lines[i])
                if m:
                    items.append(inline(m.group(2).strip()))
                    i += 1
                elif lines[i].strip() and lines[i].startswith(("  ", "\t")) and items:   # продолжение пункта
                    items[-1] += inline(" " + lines[i].strip())
                    i += 1
                else:
                    break
            blocks.append(Block("ol" if ordered else "ul", items=items))
            continue
        para.append(line)
        i += 1
    flush()
    return blocks


def document_title(markdown: str, fallback: str = "Документ") -> str:
    for block in parse(markdown):
        if block.kind == "heading" and block.level == 1:
            return plain(block.runs)
    return fallback


def column_widths(block: Block, total: float) -> list[float]:
    """Ширины колонок таблицы пропорционально содержимому: не уже самого длинного слова, но и не шире 40 знаков."""
    weights = []
    for c in range(len(block.header)):
        cells = [plain(block.header[c])] + [plain(row[c]) for row in block.rows[:80]]
        longest = max((len(w) for text in cells for w in text.split()), default=3)
        widest = max(len(t) for t in cells)
        weights.append(max(longest + 2, min(widest, 40) * 1.15 + 2, 4))
    scale = total / sum(weights)
    return [w * scale for w in weights]


# --------------------------------------------------------------------------- HTML

_CSS = """
:root { color-scheme: light; --ink: #1d1d1f; --muted: #6b6f76; --line: #d8dbe0; --accent: #2f6bff; --fill: #f3f4f6; }
* { box-sizing: border-box; }
body { margin: 0; background: #fff; color: var(--ink); font: 16px/1.6 -apple-system, "Segoe UI", "Helvetica Neue", Arial, "DejaVu Sans", sans-serif; }
main { max-width: 800px; margin: 0 auto; padding: 48px 32px 64px; }
h1 { font-size: 28px; line-height: 1.25; margin: 0 0 6px; }
h2 { font-size: 21px; margin: 32px 0 10px; padding-bottom: 6px; border-bottom: 1px solid var(--line); }
h3 { font-size: 17px; margin: 22px 0 6px; }
h4, h5, h6 { font-size: 15px; margin: 18px 0 4px; }
p { margin: 0 0 12px; } ul, ol { margin: 0 0 14px; padding-left: 26px; } li { margin: 3px 0; }
table { width: 100%; border-collapse: collapse; margin: 6px 0 18px; font-size: 14px; }
th, td { border: 1px solid var(--line); padding: 7px 10px; text-align: left; vertical-align: top; overflow-wrap: anywhere; }
th { background: var(--fill); font-weight: 600; } tr { break-inside: avoid; } thead { display: table-header-group; }
blockquote { margin: 0 0 14px; padding: 10px 14px; background: var(--fill); border-left: 4px solid var(--accent); border-radius: 4px; }
blockquote p:last-child { margin-bottom: 0; } .callout-title { font-weight: 600; margin-bottom: 4px; }
pre { margin: 0 0 14px; padding: 12px 14px; background: var(--fill); border-radius: 6px; overflow: auto; font-size: 13px; line-height: 1.45; white-space: pre-wrap; overflow-wrap: anywhere; }
code { font-family: ui-monospace, "SF Mono", Menlo, Consolas, "DejaVu Sans Mono", monospace; font-size: .92em; background: var(--fill); padding: 1px 5px; border-radius: 4px; }
pre code { background: none; padding: 0; }
a { color: var(--accent); } hr { border: none; border-top: 1px solid var(--line); margin: 24px 0; }
.doc-meta { color: var(--muted); font-size: 13px; margin-bottom: 26px; }
@page { size: A4; margin: 18mm 16mm; }
@media print { main { max-width: none; padding: 0; } h2, h3 { break-after: avoid; } pre, blockquote, table { break-inside: auto; } }
""".strip()


def _html_runs(runs: list[Run]) -> str:
    out = []
    for run in runs:
        piece = html.escape(run.text).replace("\n", "<br>")
        if run.code:
            piece = f"<code>{piece}</code>"
        if run.bold:
            piece = f"<strong>{piece}</strong>"
        if run.italic:
            piece = f"<em>{piece}</em>"
        if run.href:
            piece = f'<a href="{html.escape(run.href, quote=True)}">{piece}</a>'
        out.append(piece)
    return "".join(out)


def render_html(title: str, markdown: str, when: datetime | None = None) -> bytes:
    when = when or datetime.now()
    blocks = parse(markdown)
    body: list[str] = []
    has_h1 = any(b.kind == "heading" and b.level == 1 for b in blocks)
    if not has_h1:
        body.append(f"<h1>{html.escape(title)}</h1>")
    shown_meta = False
    for block in blocks:
        if block.kind == "heading":
            body.append(f"<h{block.level}>{_html_runs(block.runs)}</h{block.level}>")
            if block.level == 1 and not shown_meta:
                body.append(f'<div class="doc-meta">Memory Desk · {when:%d.%m.%Y}</div>')
                shown_meta = True
        elif block.kind == "para":
            body.append(f"<p>{_html_runs(block.runs)}</p>")
        elif block.kind in {"ul", "ol"}:
            body.append(f"<{block.kind}>" + "".join(f"<li>{_html_runs(item)}</li>" for item in block.items) + f"</{block.kind}>")
        elif block.kind == "table":
            head = "".join(f"<th>{_html_runs(c)}</th>" for c in block.header)
            rows = "".join("<tr>" + "".join(f"<td>{_html_runs(c)}</td>" for c in row) + "</tr>" for row in block.rows)
            body.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>")
        elif block.kind == "quote":
            inner = (f'<div class="callout-title">{html.escape(block.title)}</div>' if block.title else "")
            inner += "".join(f"<p>{_html_runs(p)}</p>" for p in block.lines)
            body.append(f"<blockquote>{inner}</blockquote>")
        elif block.kind == "code":
            body.append(f"<pre><code>{html.escape(block.text)}</code></pre>")
        elif block.kind == "hr":
            body.append("<hr>")
    page = (f'<!DOCTYPE html>\n<html lang="ru">\n<head>\n<meta charset="utf-8">\n'
            f'<meta name="viewport" content="width=device-width, initial-scale=1">\n<meta name="generator" content="Memory Desk">\n'
            f"<title>{html.escape(title)}</title>\n<style>\n{_CSS}\n</style>\n</head>\n<body>\n<main>\n"
            + "\n".join(body) + "\n</main>\n</body>\n</html>\n")
    return page.encode("utf-8")


# --------------------------------------------------------------------------- DOCX

def render_docx(title: str, markdown: str, when: datetime | None = None) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:
        raise ExportError("Не установлен модуль python-docx: запустите install.command или start.command") from exc
    when = when or datetime.now()
    blocks = parse(markdown)
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = "Memory Desk"
    doc.core_properties.language = "ru-RU"
    doc.core_properties.created = when
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2)

    def set_fonts(style, name: str, size: float | None = None) -> None:
        style.font.name = name
        if size:
            style.font.size = Pt(size)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), name)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):   # тема перекрывает имя шрифта
            if fonts.get(qn(attr)) is not None:
                del fonts.attrib[qn(attr)]

    styles = doc.styles
    set_fonts(styles["Normal"], "Calibri", 11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 22), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 11)):
        set_fonts(styles[name], "Calibri", size)
        styles[name].font.color.rgb = RGBColor(0x1D, 0x1D, 0x1F)
        styles[name].font.bold = True
    for name in ("List Bullet", "List Number", "Quote"):
        set_fonts(styles[name], "Calibri", 11)

    def add_runs(paragraph, runs: list[Run], size: float | None = None, color: RGBColor | None = None) -> None:
        for run in runs:
            if run.text == "\n":
                paragraph.add_run().add_break()
                continue
            item = paragraph.add_run(run.text)
            item.bold = run.bold or None
            item.italic = run.italic or None
            if run.code:
                item.font.name = "Consolas"
                item._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
                item.font.size = Pt((size or 11) - 1)
            elif size:
                item.font.size = Pt(size)
            if run.href:
                item.font.color.rgb = RGBColor(0x2F, 0x6B, 0xFF)
                item.font.underline = True
            elif color:
                item.font.color.rgb = color

    def shade(cell, fill: str) -> None:
        props = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        props.append(shd)

    numbering = doc.part.numbering_part.numbering_definitions._numbering
    number_style_id = styles["List Number"].element.pPr.numPr.numId.val if styles["List Number"].element.pPr is not None and styles["List Number"].element.pPr.numPr is not None else None

    def new_numbering() -> int | None:
        """Новый нумерованный список начинается с 1: у каждого списка свой `w:num` с `startOverride`."""
        if number_style_id is None:
            return None
        base = next((n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == number_style_id), None)
        abstract = base.find(qn("w:abstractNumId")).get(qn("w:val")) if base is not None else None
        if abstract is None:
            return None
        ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        num_id = max(ids) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract)
        num.append(ref)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        numbering.append(num)
        return num_id

    # колонтитулы: название и номер страницы
    header = section.header.paragraphs[0]
    header.text = title
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(0x6B, 0x6F, 0x76)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = footer.add_run("Memory Desk · стр. ")
    label.font.size = Pt(9)
    label.font.color.rgb = RGBColor(0x6B, 0x6F, 0x76)
    page = footer.add_run()
    page.font.size = Pt(9)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        page._element.append(element)

    if not any(b.kind == "heading" and b.level == 1 for b in blocks):
        doc.add_paragraph(title, style="Heading 1")
    meta_done = False
    for block in blocks:
        if block.kind == "heading":
            level = min(block.level, 4)
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            add_runs(paragraph, block.runs)
            if block.level == 1 and not meta_done:
                meta = doc.add_paragraph()
                add_runs(meta, [Run(f"Memory Desk · {when:%d.%m.%Y}")], 9, RGBColor(0x6B, 0x6F, 0x76))
                meta_done = True
        elif block.kind == "para":
            add_runs(doc.add_paragraph(), block.runs)
        elif block.kind == "ul":
            for item in block.items:
                add_runs(doc.add_paragraph(style="List Bullet"), item)
        elif block.kind == "ol":
            num_id = new_numbering()
            for item in block.items:
                paragraph = doc.add_paragraph(style="List Number")
                if num_id is not None:
                    ppr = paragraph._p.get_or_add_pPr()
                    num_pr = OxmlElement("w:numPr")
                    ilvl = OxmlElement("w:ilvl")
                    ilvl.set(qn("w:val"), "0")
                    nid = OxmlElement("w:numId")
                    nid.set(qn("w:val"), str(num_id))
                    num_pr.append(ilvl)
                    num_pr.append(nid)
                    ppr.append(num_pr)
                add_runs(paragraph, item)
        elif block.kind == "table":
            columns = len(block.header)
            table = doc.add_table(rows=1, cols=columns)
            table.style = "Table Grid"
            table.autofit = False
            for cell, runs in zip(table.rows[0].cells, block.header):
                cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                add_runs(cell.paragraphs[0], [Run(r.text, True, r.italic, r.code, r.href) for r in runs], 10)
                shade(cell, "F3F4F6")
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            tr_pr.append(repeat)
            for row in block.rows:
                cells = table.add_row().cells
                for cell, runs in zip(cells, row):
                    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                    add_runs(cell.paragraphs[0], runs, 10)
            widths = column_widths(block, 16.6)   # ширина текста страницы A4 с полями 2,2 см
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)   # сетка таблицы: по ней колонки строят LibreOffice и Pages
            for row in table.rows:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))   # строка не разрывается между страницами
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif block.kind == "quote":
            if block.title:
                head = doc.add_paragraph(style="Quote")
                add_runs(head, [Run(block.title, True)])
            for runs in block.lines:
                add_runs(doc.add_paragraph(style="Quote"), runs)
        elif block.kind == "code":
            for line in (block.text.split("\n") or [""]):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                add_runs(paragraph, [Run(line or " ", code=True)], 10)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif block.kind == "hr":
            doc.add_paragraph("―" * 20).alignment = WD_ALIGN_PARAGRAPH.CENTER
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- PDF

_fonts_ready = False


def _register_fonts() -> None:
    global _fonts_ready
    if _fonts_ready:
        return
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    files = {"DD": "DejaVuSans.ttf", "DD-Bold": "DejaVuSans-Bold.ttf", "DD-Italic": "DejaVuSans-Oblique.ttf",
             "DD-BoldItalic": "DejaVuSans-BoldOblique.ttf", "DD-Mono": "DejaVuSansMono.ttf"}
    for name, file in files.items():
        path = FONT_DIR / file
        if not path.exists():
            raise ExportError(f"Нет файла шрифта {file}: переустановите приложение")
        pdfmetrics.registerFont(TTFont(name, str(path)))
    addMapping("DD", 0, 0, "DD"), addMapping("DD", 1, 0, "DD-Bold"), addMapping("DD", 0, 1, "DD-Italic"), addMapping("DD", 1, 1, "DD-BoldItalic")
    _fonts_ready = True


_DOTS = {"🟡": "#F2B400", "🔴": "#E5484D", "🟢": "#30A46C", "🔵": "#2F6BFF", "🟣": "#8E4EC6"}


def _pdf_runs(runs: list[Run], base_bold: bool = False) -> str:
    out = []
    for run in runs:
        piece = html.escape(run.text, quote=False).replace("⏱", "◷").replace("\n", "<br/>")   # в DejaVu нет секундомера и цветных эмодзи
        for emoji, color in _DOTS.items():
            piece = piece.replace(emoji, f'<font color="{color}">●</font>')
        if run.code:
            piece = f'<font name="DD-Mono" size="8.6" backColor="#F3F4F6">{piece}</font>'
        if run.bold or base_bold:
            piece = f"<b>{piece}</b>"
        if run.italic:
            piece = f"<i>{piece}</i>"
        if run.href:
            piece = f'<a href="{html.escape(run.href, quote=True)}" color="#2F6BFF">{piece}</a>'
        out.append(piece)
    return "".join(out) or "&nbsp;"


def render_pdf(title: str, markdown: str, when: datetime | None = None) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import (BaseDocTemplate, Flowable, Frame, KeepTogether, ListFlowable, ListItem, PageTemplate,
                                        Paragraph, Spacer, Table, TableStyle)
    except ImportError as exc:
        raise ExportError("Не установлен модуль reportlab: запустите install.command или start.command") from exc
    _register_fonts()
    when = when or datetime.now()
    blocks = parse(markdown)
    ink, muted, line_color, fill = colors.HexColor("#1D1D1F"), colors.HexColor("#6B6F76"), colors.HexColor("#D8DBE0"), colors.HexColor("#F3F4F6")
    base = ParagraphStyle("base", fontName="DD", fontSize=10, leading=14.5, textColor=ink, alignment=TA_LEFT, spaceAfter=6)
    heads = {1: ParagraphStyle("h1", parent=base, fontName="DD-Bold", fontSize=19, leading=24, spaceAfter=3),
             2: ParagraphStyle("h2", parent=base, fontName="DD-Bold", fontSize=14, leading=18, spaceBefore=14, spaceAfter=6),
             3: ParagraphStyle("h3", parent=base, fontName="DD-Bold", fontSize=11.5, leading=15, spaceBefore=10, spaceAfter=4)}
    heads[4] = heads[5] = heads[6] = ParagraphStyle("h4", parent=heads[3], fontSize=10.5)
    meta_style = ParagraphStyle("meta", parent=base, fontSize=8.5, leading=11, textColor=muted, spaceAfter=14)
    cell_style = ParagraphStyle("cell", parent=base, fontSize=9, leading=12, spaceAfter=0)
    head_cell = ParagraphStyle("hcell", parent=cell_style, fontName="DD-Bold")
    quote_style = ParagraphStyle("quote", parent=base, leftIndent=8, spaceAfter=3)
    code_style = ParagraphStyle("code", parent=base, fontName="DD-Mono", fontSize=8, leading=10.5, spaceAfter=0)
    page_width = A4[0] - 36 * mm

    class Heading(Flowable):
        """Заголовок с закладкой в оглавлении PDF."""
        def __init__(self, paragraph, level, key):
            super().__init__()
            self.p, self.level, self.key = paragraph, level, key

        def wrap(self, w, h):
            width, height = self.p.wrap(w, h)
            self.width, self.height = width, height
            return width, height

        def split(self, w, h):
            return []

        def draw(self):
            self.canv.bookmarkPage(self.key)
            if self.level <= 2:
                self.canv.addOutlineEntry(re.sub(r"<[^>]+>", "", self.p.text) or "…", self.key, level=self.level - 1, closed=False)
            self.p.drawOn(self.canv, 0, 0)

    story: list = []
    counter = 0
    has_h1 = any(b.kind == "heading" and b.level == 1 for b in blocks)
    if not has_h1:
        story += [Paragraph(html.escape(title), heads[1]), Paragraph(f"Memory Desk · {when:%d.%m.%Y}", meta_style)]
    meta_done = has_h1 is False
    for block in blocks:
        if block.kind == "heading":
            counter += 1
            story.append(Heading(Paragraph(_pdf_runs(block.runs), heads[block.level]), min(block.level, 3), f"h{counter}"))
            if block.level == 1 and not meta_done:
                story.append(Paragraph(f"Memory Desk · {when:%d.%m.%Y}", meta_style))
                meta_done = True
        elif block.kind == "para":
            story.append(Paragraph(_pdf_runs(block.runs), base))
        elif block.kind in {"ul", "ol"}:
            items = [ListItem(Paragraph(_pdf_runs(item), ParagraphStyle("li", parent=base, spaceAfter=2)), leftIndent=16) for item in block.items]
            story.append(ListFlowable(items, bulletType="bullet" if block.kind == "ul" else "1", start="•" if block.kind == "ul" else 1,
                                      bulletFontName="DD", bulletFontSize=9.5, bulletFormat=None if block.kind == "ul" else "%s.", leftIndent=16, bulletDedent=14, spaceAfter=6))
        elif block.kind == "table":
            columns = len(block.header)
            from reportlab.pdfbase.pdfmetrics import stringWidth
            widths = [w * mm for w in column_widths(block, page_width / mm)]
            floor = []   # ни одно слово не должно ломаться посреди: колонка не уже самого широкого слова (шапка жирная)
            for c in range(columns):
                words = [(w, "DD-Bold") for w in plain(block.header[c]).split()]
                words += [(w, "DD") for row in block.rows[:80] for w in plain(row[c]).split()]
                floor.append(max((stringWidth(w, f, 9) for w, f in words), default=10) + 11)
            fixed = {c for c in range(columns) if widths[c] < floor[c]}
            for c in fixed:
                widths[c] = floor[c]
            spare = page_width - sum(widths[c] for c in fixed)
            rest = sum(widths[c] for c in range(columns) if c not in fixed)
            for c in range(columns):
                if c not in fixed and rest:
                    widths[c] = max(floor[c], widths[c] * spare / rest)
            data = [[Paragraph(_pdf_runs(c), head_cell) for c in block.header]] + [[Paragraph(_pdf_runs(c), cell_style) for c in row] for row in block.rows]
            table = Table(data, colWidths=widths, repeatRows=1, splitByRow=1)
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, line_color), ("BACKGROUND", (0, 0), (-1, 0), fill),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                                       ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
            story += [table, Spacer(1, 8)]
        elif block.kind == "quote":
            rows = []
            if block.title:
                rows.append([Paragraph(f"<b>{html.escape(block.title)}</b>", quote_style)])
            rows += [[Paragraph(_pdf_runs(p), quote_style)] for p in block.lines]
            if rows:
                box = Table(rows, colWidths=[page_width])
                box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), fill), ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor("#2F6BFF")),
                                         ("LEFTPADDING", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
                story += [box, Spacer(1, 8)]
        elif block.kind == "code":
            rows = [[Paragraph(html.escape(l).replace(" ", "&nbsp;") or "&nbsp;", code_style)] for l in (block.text.split("\n") or [""])]
            box = Table(rows, colWidths=[page_width])
            box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), fill), ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 0),
                                     ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
            story += [box, Spacer(1, 8)]
        elif block.kind == "hr":
            rule = Table([[""]], colWidths=[page_width], rowHeights=[1])
            rule.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.6, line_color)]))
            story += [Spacer(1, 6), rule, Spacer(1, 8)]
    if not story:
        story.append(Paragraph("Документ пуст.", base))

    def decorate(canvas, doc_) -> None:
        canvas.saveState()
        canvas.setFont("DD", 8)
        canvas.setFillColor(muted)
        canvas.drawString(18 * mm, A4[1] - 11 * mm, (title[:90] + "…") if len(title) > 90 else title)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Memory Desk · стр. {doc_.page}")
        canvas.restoreState()

    out = io.BytesIO()
    doc = BaseDocTemplate(out, pagesize=A4, title=title, author="Memory Desk", leftMargin=18 * mm, rightMargin=18 * mm,
                          topMargin=18 * mm, bottomMargin=18 * mm, lang="ru-RU")
    doc.addPageTemplates([PageTemplate(id="main", frames=[Frame(18 * mm, 18 * mm, page_width, A4[1] - 36 * mm, id="body", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)],
                                       onPage=decorate)])
    try:
        doc.build(story)
    except Exception as exc:  # noqa: BLE001 — reportlab бросает разные ошибки вёрстки
        raise ExportError(f"Не удалось собрать PDF: {exc}") from exc
    return out.getvalue()


# --------------------------------------------------------------------------- единая точка входа

def safe_filename(name: str, limit: int = 90) -> str:
    name = re.sub(r":\s+", " — ", " ".join(str(name).split()))
    cleaned = re.sub(r"[\\/:*?\"<>|\x00-\x1f]", "-", name).strip(" .-")
    return (cleaned[:limit].rstrip(" .-,;") or "Документ")


def content_disposition(filename: str) -> str:
    ascii_name = re.sub(r"[^\x20-\x7e]", "_", filename).replace('"', "")
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"


def export(fmt: str, title: str, markdown: str, name: str | None = None, when: datetime | None = None) -> tuple[bytes, str, str]:
    """Возвращает (содержимое, имя файла, тип). `name` — имя без расширения (по умолчанию из названия документа)."""
    if fmt not in FORMATS:
        raise ExportError("Неизвестный формат: нужен md, html, pdf или docx")
    if not markdown.strip():
        raise ExportError("Нечего выгружать: документ пуст")
    if len(markdown) > MAX_CHARS:
        raise ExportError("Документ слишком большой для выгрузки")
    title = " ".join((title or "").split()) or document_title(markdown)
    stem = safe_filename(re.sub(r"\.(md|html|pdf|docx)$", "", name or title, flags=re.I))
    if fmt == "md":
        data = (markdown.rstrip() + "\n").encode("utf-8")
    elif fmt == "html":
        data = render_html(title, markdown, when)
    elif fmt == "docx":
        data = render_docx(title, markdown, when)
    else:
        data = render_pdf(title, markdown, when)
    return data, f"{stem}.{fmt}", FORMATS[fmt][0]
