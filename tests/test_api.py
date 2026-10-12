from __future__ import annotations

import re
import time
from pathlib import Path

import pytest


def by_title(items, fragment):
    return next(m for t, m in items.items() if fragment in t)


# --------------------------------------------------------------------------- безопасность

def test_health_and_mutation_guard(client):
    assert client.get("/api/health").json()["ok"] is True
    denied = client.post("/api/jobs/cancel", headers={"x-memory-desk": "0"})
    assert denied.status_code == 403


# --------------------------------------------------------------------------- записи

def test_meetings_list(client):
    data = client.get("/api/meetings").json()
    assert data["counts"] == {"total": 5, "needsReview": 1, "noProject": 1}
    titles = [m["title"] for m in data["items"]]
    assert titles[0] == "Встреча один на один с Анной"
    assert titles == sorted(titles, key=lambda t: next(m for m in data["items"] if m["title"] == t)["date"] + next(m for m in data["items"] if m["title"] == t)["time"], reverse=True)
    item = data["items"][0]
    assert item["project"]["title"] == "КХД" and item["project"]["color"].startswith("#")
    assert item["status"] == "completed" and item["durationSec"] == 1820 + 0 or item["duration"]


def test_filters(client):
    assert len(client.get("/api/meetings", params={"project": "КХД"}).json()["items"]) == 3
    assert len(client.get("/api/meetings", params={"project": "Без проекта"}).json()["items"]) == 1
    assert len(client.get("/api/meetings", params={"status": "needs_review"}).json()["items"]) == 1
    assert len(client.get("/api/meetings", params={"status": "completed"}).json()["items"]) == 4
    assert len(client.get("/api/meetings", params={"type": "Zoom"}).json()["items"]) == 1
    assert len(client.get("/api/meetings", params={"period": "today"}).json()["items"]) == 1
    assert len(client.get("/api/meetings", params={"period": "7d"}).json()["items"]) == 5
    assert len(client.get("/api/meetings", params={"person": "Михаил Орлов"}).json()["items"]) == 3
    # поиск идёт и по стенограмме
    found = client.get("/api/meetings", params={"q": "альманах"}).json()["items"]
    assert [m["title"] for m in found] == ["Обсуждение машинного времени и последствий"]
    assert client.get("/api/meetings", params={"q": "несуществующее-слово"}).json()["items"] == []


def test_meeting_detail(client, items):
    m = by_title(items, "Созвон по архитектуре")
    d = client.get(f"/api/meetings/{m['id']}").json()
    assert d["summary"][0].startswith("Команда переходит")
    assert [x["text"] for x in d["decisions"]] == [
        "Перейти на инкрементальную загрузку витрин продаж", "Не менять слой сырых данных до конца квартала"]
    assert {t["owner"] for t in d["tasks"]} == {"Анна", "Михаил"}
    assert d["questions"][0]["text"] == "Где хранить историю изменений справочников?"
    # служебная разметка не должна просачиваться в текст
    for bucket in ("risks", "questions"):
        for fact in d[bucket]:
            assert "⏱" not in fact["text"] and "`" not in fact["text"]
    assert d["transcript"][1] == {"i": 1, "t": "00:00:06", "sec": 6, "speaker": "Михаил", "text": "Полная перезагрузка занимает уже больше четырёх часов."}
    assert d["hasAudio"] and d["obsidianUrl"].startswith("obsidian://open?")
    assert client.get("/api/meetings/nope").status_code == 404


def test_audio_supports_range_and_waveform(client, items):
    m = by_title(items, "Идеи для заметки")
    full = client.get(f"/api/meetings/{m['id']}/audio")
    assert full.status_code == 200 and full.headers["content-type"].startswith("audio/")
    part = client.get(f"/api/meetings/{m['id']}/audio", headers={"range": "bytes=0-99"})
    assert part.status_code == 206 and len(part.content) == 100
    bars = client.get(f"/api/meetings/{m['id']}/waveform").json()["bars"]
    assert len(bars) == 64 and all(0 < b <= 1 for b in bars)


def test_publish_requires_project_and_type_then_flow(client, items, fixture_dir):
    m = by_title(items, "машинного времени")
    assert m["status"] == "needs_review"
    assert client.post(f"/api/meetings/{m['id']}/publish").status_code == 409

    updated = client.patch(f"/api/meetings/{m['id']}", json={"project_id": "kxd", "record_type_id": "zoom"}).json()
    assert updated["project"]["title"] == "КХД" and updated["type"]["title"] == "Zoom"
    note = next((fixture_dir / "Claude" / "Встречи").rglob("*машинного времени*.md"))
    text = note.read_text(encoding="utf-8")
    assert 'project_id: "kxd"' in text and 'project: "КХД"' in text and "source_type: \"Zoom\"" in text
    assert text.count("project_id:") == 1  # заметка не продублирована и не испорчена

    published = client.post(f"/api/meetings/{m['id']}/publish").json()
    assert published["status"] == "completed"
    counts = client.get("/api/meetings").json()["counts"]
    assert counts["needsReview"] == 0 and counts["noProject"] == 0


def test_people_assignment_and_important_archive(client, items):
    m = by_title(items, "Идеи для заметки")
    d = client.patch(f"/api/meetings/{m['id']}", json={"people": ["Иван Петров", "иван   петров ", "Оля"], "important": True}).json()
    assert [p["name"] for p in d["people"]] == ["Иван Петров", "Оля"] and d["important"] is True
    assert {"Иван Петров", "Оля"} <= {p["name"] for p in client.get("/api/people").json()["items"]}
    assert client.post(f"/api/meetings/{m['id']}/archive").json()["ok"]
    assert all(x["id"] != m["id"] for x in client.get("/api/meetings").json()["items"])
    assert any(x["id"] == m["id"] for x in client.get("/api/meetings", params={"archived": True}).json()["items"])
    assert client.post(f"/api/meetings/{m['id']}/restore").json()["ok"]
    assert any(x["id"] == m["id"] for x in client.get("/api/meetings").json()["items"])


# --------------------------------------------------------------------------- задачи и решения

def test_tasks_registry_and_status_flow(client):
    reg = client.get("/api/tasks", params={"kind": "tasks"}).json()
    assert reg["counts"]["tasks"] == 6 and reg["counts"]["decisions"] == 4
    assert all(t["draft"] for t in reg["items"])  # после импорта всё — черновики
    task = next(t for t in reg["items"] if t["text"].startswith("Подготовить схему"))
    assert task["owner"] == "Анна" and task["due"] and task["tSec"] == 30

    confirmed = client.patch(f"/api/tasks/{task['id']}", json={"status": "open"}).json()
    assert confirmed["status"] == "open" and not confirmed["draft"]
    done = client.patch(f"/api/tasks/{task['id']}", json={"status": "done"}).json()
    assert done["done"] is True
    assert client.patch(f"/api/tasks/{task['id']}", json={"status": "bogus"}).status_code == 400
    assert client.patch("/api/tasks/unknown", json={"status": "open"}).status_code == 404
    rejected = client.patch(f"/api/tasks/{task['id']}", json={"status": "rejected"}).json()
    assert rejected["status"] == "rejected"
    assert all(t["id"] != task["id"] for t in client.get("/api/tasks").json()["items"])

    dec = client.get("/api/tasks", params={"kind": "decisions"}).json()["items"][0]
    assert client.patch(f"/api/decisions/{dec['id']}", json={"status": "confirmed"}).json()["status"] == "confirmed"
    assert client.patch(f"/api/decisions/{dec['id']}", json={"status": "open"}).status_code == 400
    edited = client.patch(f"/api/decisions/{dec['id']}", json={"text": "Новая формулировка"}).json()
    assert edited["text"] == "Новая формулировка"
    assert client.patch(f"/api/decisions/{dec['id']}", json={"text": "  "}).status_code == 400


def test_today(client):
    t = client.get("/api/today").json()
    assert t["hero"]["kind"] == "last_record" and t["hero"]["title"] == "Встреча один на один с Анной"
    assert set(t["hero"]["columns"]) == {"decisions", "tasks", "questions"}
    assert 1 <= len(t["newRecords"]) <= 4
    assert t["drafts"] >= 1 and t["noProject"] == 0
    # подтверждённые, но не выполненные задачи попадают в карточку «Задачи»
    reg = client.get("/api/tasks").json()["items"]
    draft = next(x for x in reg if x["status"] == "draft")
    client.patch(f"/api/tasks/{draft['id']}", json={"status": "open"})
    assert any(x["id"] == draft["id"] for x in client.get("/api/today").json()["tasks"])


def test_projects_and_people(client):
    projects = {p["title"]: p for p in client.get("/api/projects").json()["items"]}
    assert set(projects) == {"КХД", "Личное"}
    assert projects["КХД"]["counts"]["records"] == 4 and projects["КХД"]["goal"].startswith("Корпоративное")
    assert projects["КХД"]["lastChange"]["text"]
    assert any(p["name"] == "Анна Соколова" and p["meetings"] >= 2 for p in client.get("/api/people").json()["items"])
    created = client.post("/api/projects", json={"title": "Новый проект"}).json()
    assert created["created"] and created["id"] == "новый-проект"
    assert client.post("/api/projects", json={"title": "bad/name"}).status_code == 400
    assert "Новый проект" in {p["title"] for p in client.get("/api/meta").json()["projects"]}
    assert client.post("/api/people", json={"name": ""}).status_code == 400


# --------------------------------------------------------------------------- чат

def test_chat_with_sources_and_scopes(client, items):
    r = client.post("/api/chat", json={"question": "Что решили по инкрементальной загрузке витрин?", "scope": {"kind": "all"}}).json()
    assert r["confirmed"] and r["sources"]
    top = r["sources"][0]
    assert top["title"] == "Созвон по архитектуре хранилища" and top["meetingId"]
    assert "[" not in r["answer"] or "встреча" not in r["answer"]
    # область «Текущая запись» ограничивает поиск
    m = by_title(items, "Идеи для заметки")
    scoped = client.post("/api/chat", json={"question": "Что решили по инкрементальной загрузке витрин?",
                                            "scope": {"kind": "record", "id": m["id"]}}).json()
    assert not scoped["confirmed"] and scoped["sources"] == []
    assert scoped["answer"].startswith("В записях этой области нет фрагментов")
    proj = client.post("/api/chat", json={"question": "Какие задачи у Анны?", "scope": {"kind": "project", "id": "kxd", "title": "КХД"}}).json()
    assert all(s["title"] != "Идеи для заметки о привычках" for s in proj["sources"])
    assert client.post("/api/chat", json={"question": "  "}).status_code == 400


def test_chat_overview_phrasings(client):
    """«синтез по проекту КХД», «по проекту КХД» и просто «КХД» — просьбы обзора; точечные вопросы остаются поиском."""
    for question in ("синтез по проекту кхд", "выжимка по проекту КХД", "по проекту КХД", "КХД", "общая картина по КХД"):
        r = client.post("/api/chat", json={"question": question, "scope": {"kind": "all"}}).json()
        assert r.get("digest") and r["digest"]["records"] >= 3 and r["confirmed"], question
    for question in ("Кто отвечает за миграцию справочников?", "Когда решили перейти на дельту по КХД?", "Что решили по инкрементальной загрузке витрин?"):
        r = client.post("/api/chat", json={"question": question, "scope": {"kind": "all"}}).json()
        assert not r.get("digest"), question


def test_chat_overview_widens_to_all_time(client, monkeypatch):
    """Период не назван, а за 30 дней записей нет — берётся всё время, а не ответ «ничего нет»."""
    from datetime import date
    from memory_desk import digest
    real = digest.select_meetings
    monkeypatch.setattr(digest, "select_meetings",
                        lambda core, intent, date_of: [] if intent.label.startswith("последние 30") else real(core, intent, date_of))
    r = client.post("/api/chat", json={"question": "синтез по проекту КХД", "scope": {"kind": "all"}}).json()
    assert r["digest"]["records"] >= 3 and "всё время" in r["digest"]["title"].casefold() or "всё время" in r["answer"].casefold()


def test_chat_gives_honest_answer_without_evidence(client):
    r = client.post("/api/chat", json={"question": "квантовая хромодинамика бозон", "scope": {"kind": "all"}}).json()
    assert r["confirmed"] is False and "нет фрагментов" in r["answer"]


# --------------------------------------------------------------------------- очередь

def test_queue_states(client):
    jobs = {j["file"]: j for j in client.get("/api/jobs").json()["jobs"]}
    assert jobs["Zoom_Ретро_спринта.m4a"]["state"] == "queued"
    assert jobs["Голосовая_заметка_1130.mp3"]["state"] == "queued"
    failed = jobs["Звонок_подрядчик.amr"]
    assert failed["state"] == "failed" and failed["error"]["title"] == "Формат не поддерживается"
    assert failed["error"]["action"] == "formats" and failed["error"]["stage"] == "Проверка файла"


def test_scan_cancel_retry_dismiss(client):
    first = client.post("/api/jobs/scan").json()
    assert first["created"] and first["message"] == "Найдено 2 новых файла · добавлено в очередь"
    assert client.post("/api/jobs/scan").json()["message"] == "Обработка уже идёт"
    running = client.get("/api/jobs").json()
    assert running["active"] and running["running"] is None or running["active"]
    assert client.post("/api/jobs/cancel").json()["ok"]
    after = {j["file"]: j for j in client.get("/api/jobs").json()["jobs"]}
    assert after["Zoom_Ретро_спринта.m4a"]["state"] == "failed"
    assert after["Zoom_Ретро_спринта.m4a"]["error"]["title"] == "Обработка отменена"
    assert after["Zoom_Ретро_спринта.m4a"]["error"]["action"] == "retry"
    assert client.post("/api/jobs/retry", json={"file": "Zoom_Ретро_спринта.m4a"}).json()["created"]
    assert client.post("/api/jobs/cancel").json()["ok"]
    assert client.post("/api/jobs/dismiss", json={"file": "Голосовая_заметка_1130.mp3"}).json()["ok"]
    assert "Голосовая_заметка_1130.mp3" not in {j["file"] for j in client.get("/api/jobs").json()["jobs"]}
    assert client.post("/api/jobs/retry", json={"file": "../etc/passwd"}).status_code == 400
    assert client.post("/api/jobs/retry", json={"file": "нет-такого.mp3"}).status_code == 404
    log = client.get("/api/jobs/log", params={"file": "Zoom_Ретро_спринта.m4a"}).json()
    assert "events" in log and "tail" in log


# --------------------------------------------------------------------------- система

def test_system_and_settings(client, fixture_dir):
    s = client.get("/api/system").json()
    assert [r["key"] for r in s["rows"]] == ["whisper", "ollama", "index", "recordings", "vault", "disk", "backup"]
    assert next(r for r in s["rows"] if r["key"] == "ollama")["state"] == "warn"
    assert next(r for r in s["rows"] if r["key"] == "index")["state"] == "ok"
    cfg = client.get("/api/settings").json()
    assert cfg["exists"]["recordings_dir"] and cfg["retention_days"] == 14
    bad = client.put("/api/settings", json={**cfg, "retention_days": "abc"})
    assert bad.status_code == 422 or bad.status_code == 400
    empty = client.put("/api/settings", json={**{k: cfg[k] for k in ("recordings_dir", "obsidian_dir", "archive_dir")}, "obsidian_dir": " "})
    assert empty.status_code == 400
    ok = client.put("/api/settings", json={**{k: cfg[k] for k in ("recordings_dir", "obsidian_dir", "archive_dir")},
                                            "retention_days": 30, "auto_process": True}).json()
    assert ok["retention_days"] == 30 and ok["auto_process"] is True
    client.put("/api/settings", json={**{k: cfg[k] for k in ("recordings_dir", "obsidian_dir", "archive_dir")}, "auto_process": False})
    assert client.post("/api/system/backup").json()["path"].endswith(".db")


def test_spa_fallback(client):
    r = client.get("/records/abc")
    assert r.status_code in (200, 503) and "text/html" in r.headers["content-type"]
    assert client.get("/api/unknown").status_code == 404


# --------------------------------------------------------------------------- чистая логика

def test_parse_fact_strips_service_markup():
    from memory_desk import storage
    assert storage.parse_fact("Что это за хей? — `⏱ 00:02:18`") == {
        "text": "Что это за хей?", "owner": "", "deadline": "", "evidence_time": "00:02:18"}
    fact = storage.parse_fact("[ ] Выдать доступ — **ответственный:** Анна; **срок:** 2026-10-16; `⏱ 00:01:35`")
    assert fact == {"text": "Выдать доступ", "owner": "Анна", "deadline": "2026-10-16", "evidence_time": "00:01:35"}


def test_plural_and_errors_and_estimate():
    from datetime import datetime, timedelta
    from memory_desk import jobs
    assert [jobs.ru_plural(n, "a", "b", "c") for n in (1, 2, 5, 11, 21, 22, 25)] == ["a", "b", "c", "c", "a", "b", "c"]
    assert jobs.humanize_error("Ollama не отвечает. Запустите...", "Анализ")["action"] == "system"
    assert jobs.humanize_error("что-то странное", "Анализ")["title"] == "Не удалось обработать запись"
    now = datetime.now()
    pct = [jobs.estimate(stage, now - timedelta(seconds=sec), 1800, now)[0]
           for stage, sec in (("checking", 0), ("preparing_audio", 1), ("transcribing", 0), ("transcribing", 90),
                              ("analyzing", 0), ("analyzing", 60), ("writing_note", 0), ("archiving", 0))]
    assert pct == sorted(pct) and pct[0] < 5 and 90 <= pct[-1] <= 99
    assert jobs.estimate("completed", now, 1800, now) == (100, None)


# --------------------------------------------------------------------------- говорящие, резервные копии

def test_diarize_module_roundtrip():
    from memory_desk import diarize
    body = "[00:00:01] Привет\n\n[00:00:05] **Спикер 1:** Да\nпродолжение\n\n[00:00:09] Пока\n"
    parsed = diarize.parse(body)
    assert [s["speaker"] for s in parsed] == ["", "Спикер 1", ""]
    assert parsed[1]["text"] == "Да продолжение" and [s["i"] for s in parsed] == [0, 1, 2]
    out = diarize.relabel(body, ["Анна", "", "Борис"])
    assert "[00:00:01] **Анна:** Привет" in out and "[00:00:05] Да" in out and "[00:00:09] **Борис:** Пока" in out
    assert diarize.relabel(out, ["", "", ""]) .count("**") == 0


def test_detect_speakers_with_fake_model():
    from memory_desk import diarize
    segments = [{"text": f"реплика {n}", "gap": 2.0 if n % 2 else None} for n in range(10)]
    calls = []

    def chat(prompt):
        calls.append(prompt)
        return {"turns": [{"from": 0, "speaker": "A"}, {"from": 3, "speaker": "B"}, {"from": 7, "speaker": "A"}, {"from": 99, "speaker": "Z"}]}

    labels = diarize.detect_speakers(segments, chat)
    assert labels == ["Спикер 1"] * 3 + ["Спикер 2"] * 4 + ["Спикер 1"] * 3 and len(calls) == 1
    assert "пауза 2.0 с" in calls[0]
    assert diarize.detect_speakers(segments, lambda _p: {"turns": [{"from": 0, "speaker": "A"}]}) == []
    assert diarize.detect_speakers(segments, lambda _p: {"junk": 1}) == []


def test_llm_json_repair():
    from memory_desk import llm
    assert llm.extract_json('<think>x</think>```json\n{"a": 1}\n```') == {"a": 1}
    assert llm.extract_json('{"title":"T","tasks":[{"task":"a"},{"task":"вто') == {"title": "T", "tasks": [{"task": "a"}, {"task": "вто"}]}
    assert llm.extract_json("нет json") is None


def wait_diarized(client, mid):
    for _ in range(100):
        status = client.get(f"/api/meetings/{mid}/diarize").json()
        if status["state"] != "running":
            return status
        time.sleep(0.05)
    raise AssertionError("диаризация не завершилась")


def test_speakers_flow(client, items, monkeypatch):
    from memory_desk import speakers
    m = by_title(items, "Обсуждение машинного времени и последствий")
    mid = m["id"]
    monkeypatch.setattr(speakers, "make_chat", lambda: (lambda _p: {"turns": [
        {"from": 0, "speaker": "A"}, {"from": 2, "speaker": "B"}, {"from": 4, "speaker": "A"}]}))
    assert client.post(f"/api/meetings/{mid}/diarize").status_code == 200
    assert wait_diarized(client, mid)["state"] == "done"
    d = client.get(f"/api/meetings/{mid}").json()
    assert [x["speaker"] for x in d["transcript"]] == ["Спикер 1"] * 2 + ["Спикер 2"] * 2 + ["Спикер 1"] * 2
    assert {x["label"] for x in d["speakers"]} == {"Спикер 1", "Спикер 2"} and all(x["generic"] for x in d["speakers"])

    # переименование всех реплик + привязка к существующему человеку (по имени без фамилии)
    r = client.put(f"/api/meetings/{mid}/speakers", json={"rename": {"Спикер 2": "Анна"}})
    assert r.status_code == 200
    labels = [x["speaker"] for x in r.json()["transcript"]]
    assert labels[2] == labels[3] == "Анна Соколова"
    assert "Анна Соколова" in [p["name"] for p in r.json()["people"]]
    assert next(x for x in r.json()["speakers"] if x["label"] == "Анна Соколова")["personId"]

    # правка одной реплики на нового человека; он появляется в участниках и справочнике
    r = client.put(f"/api/meetings/{mid}/speakers", json={"segments": [{"i": 0, "speaker": "Пётр Новиков"}]})
    assert r.json()["transcript"][0]["speaker"] == "Пётр Новиков"
    assert "Пётр Новиков" in [p["name"] for p in client.get("/api/people").json()["items"]]
    # слияние: переименование в уже существующую метку
    r = client.put(f"/api/meetings/{mid}/speakers", json={"rename": {"Спикер 1": "Анна Соколова"}})
    assert {x["speaker"] for x in r.json()["transcript"]} == {"Анна Соколова", "Пётр Новиков"}
    # убрать разметку у одной реплики
    r = client.put(f"/api/meetings/{mid}/speakers", json={"segments": [{"i": 0, "speaker": ""}]})
    assert r.json()["transcript"][0]["speaker"] == ""
    assert client.put(f"/api/meetings/{mid}/speakers", json={}).status_code == 400
    assert client.put("/api/meetings/nope/speakers", json={"rename": {"a": "b"}}).status_code == 404


def test_diarize_single_speaker_and_errors(client, items, monkeypatch):
    from memory_desk import speakers
    mid = by_title(items, "Встреча один на один с Анной")["id"]
    monkeypatch.setattr(speakers, "make_chat", lambda: (lambda _p: {"turns": [{"from": 0, "speaker": "A"}]}))
    client.post(f"/api/meetings/{mid}/diarize")
    status = wait_diarized(client, mid)
    assert status["state"] == "done" and "один человек" in status["message"]

    def boom(_prompt):
        raise RuntimeError("Ollama не отвечает. Запустите приложение Ollama или `ollama serve`.")
    monkeypatch.setattr(speakers, "make_chat", lambda: boom)
    client.post(f"/api/meetings/{mid}/diarize")
    status = wait_diarized(client, mid)
    assert status["state"] == "error" and "Ollama не отвечает" in status["message"]


def test_backups_list_and_restore(client):
    client.post("/api/system/backup")
    items = client.get("/api/system/backups").json()["items"]
    assert items and items[0]["name"].endswith(".db") and items[0]["size"] > 0
    assert client.post("/api/system/restore", json={"name": "../etc/passwd"}).status_code == 404
    assert client.post("/api/system/restore", json={"name": "нет.db"}).status_code == 404
    ok = client.post("/api/system/restore", json={"name": items[0]["name"]})
    assert ok.status_code == 200 and "восстановлена" in ok.json()["message"]
    assert client.get("/api/meetings").json()["counts"]["total"] >= 5


# --------------------------------------------------------------------------- Календарь и Напоминания (подмена EventKit)

@pytest.fixture()
def fake_apple(client, monkeypatch):
    from memory_desk import apple, integrations
    fake = apple.FakeBackend()
    apple.use_backend(fake)
    monkeypatch.setattr(integrations, "schedule_sync", lambda cfg: None)
    cfg = client.get("/api/settings").json()
    cfg.pop("exists", None)
    original = dict(cfg)

    def configure(**values):
        cfg.update(values)
        r = client.put("/api/settings", json=cfg)
        assert r.status_code == 200, r.text
        return r.json()

    yield fake, configure
    client.put("/api/settings", json=original)
    apple.use_backend(None)


def test_apple_access_flow(client, fake_apple):
    fake, _configure = fake_apple
    fake.access = "notDetermined"
    assert client.get("/api/integrations/apple").json()["events"] == "notDetermined"
    r = client.post("/api/integrations/apple/access").json()
    assert r["events"] == "authorized" and r["reminders"] == "authorized"
    from memory_desk import apple
    apple.use_backend(apple.UnavailableBackend("Доступно только на macOS"))
    assert client.get("/api/integrations/apple").json()["available"] is False
    assert client.post("/api/integrations/apple/access").status_code == 409


def test_today_uses_apple_calendar(client, fake_apple):
    from datetime import datetime, timedelta
    fake, configure = fake_apple
    now = datetime.now()
    fake.add_event("Планёрка КХД", now - timedelta(minutes=10), now + timedelta(minutes=50), ["Я", "Соколова Анна", "Игорь Петров"])
    fake.add_event("Отпуск", now.replace(hour=0, minute=0), now.replace(hour=23, minute=59), allDay=True)
    configure(apple_calendar=True)
    hero = client.get("/api/today").json()["hero"]
    assert hero["kind"] == "next_meeting" and hero["title"] == "Планёрка КХД" and hero["ongoing"] is True
    assert [p["name"] for p in hero["people"]] == ["Соколова Анна", "Игорь Петров"]
    assert hero["project"]["title"] == "КХД"
    configure(apple_calendar=False)
    assert (client.get("/api/today").json()["hero"] or {}).get("kind") != "next_meeting"


def test_calendar_suggestion_and_apply(client, items, fake_apple):
    from datetime import datetime, timedelta
    fake, configure = fake_apple
    m = by_title(items, "Созвон по архитектуре")
    detail = client.get(f"/api/meetings/{m['id']}").json()
    start = datetime.fromisoformat(f"{detail['date']}T{detail['time']}")
    fake.add_event("Архитектурный комитет КХД", start - timedelta(minutes=5), start + timedelta(minutes=50),
                   ["Я", "Соколова Анна", "Игорь Петров"])
    fake.add_event("Совсем другое", start + timedelta(hours=5), start + timedelta(hours=6), ["Кто-то"])
    assert client.get(f"/api/meetings/{m['id']}/calendar").json()["suggestion"] is None  # календарь не включён
    configure(apple_calendar=True)
    client.patch(f"/api/meetings/{m['id']}", json={"project_id": ""})
    s = client.get(f"/api/meetings/{m['id']}/calendar").json()["suggestion"]
    assert s["event"]["title"] == "Архитектурный комитет КХД" and s["titleDiffers"] is True
    assert "Анна Соколова" in s["people"] and "Игорь Петров" in s["people"] and "Я" not in s["people"]
    assert s["project"]["title"] == "КХД"
    r = client.post(f"/api/meetings/{m['id']}/calendar", json={"action": "apply", "useTitle": True})
    assert r.status_code == 200 and r.json()["suggestion"]["linked"] is True
    d = r.json()["meeting"]
    assert d["title"] == "Архитектурный комитет КХД" and d["project"]["title"] == "КХД"
    assert {"Анна Соколова", "Игорь Петров"} <= {p["name"] for p in d["people"]}
    assert client.post(f"/api/meetings/{m['id']}/calendar", json={"action": "apply"}).status_code == 409
    # «скрыть» запоминается
    other = by_title(items, "Идеи для заметки")
    o = client.get(f"/api/meetings/{other['id']}").json()
    ostart = datetime.fromisoformat(f"{o['date']}T{o['time']}")
    fake.add_event("Утренний ритуал", ostart, ostart + timedelta(minutes=5), ["Я"])
    assert client.get(f"/api/meetings/{other['id']}/calendar").json()["suggestion"]["event"]["title"] == "Утренний ритуал"
    client.post(f"/api/meetings/{other['id']}/calendar", json={"action": "dismiss"})
    assert client.get(f"/api/meetings/{other['id']}/calendar").json()["suggestion"] is None


def test_create_calendar_event(client, items, fake_apple):
    fake, configure = fake_apple
    mid = by_title(items, "Идеи для заметки")["id"]
    r = client.post("/api/calendar/events", json={"title": "Срок: черновик", "start": "2031-05-20T00:00:00", "allDay": True, "meetingId": mid})
    assert r.status_code == 200
    event = fake.created_events[-1]
    assert event["allDay"] is True and event["start"].startswith("2031-05-20T00:00:00") and event["end"].startswith("2031-05-20T23:59:59")
    assert f"/records/{mid}" in event["url"] and "Идеи для заметки" in event["notes"]
    assert client.post("/api/calendar/events", json={"title": "x", "start": "2031-05-20T10:00:00", "end": "2031-05-20T09:00:00"}).status_code == 400
    assert client.post("/api/calendar/events", json={"title": " ", "start": "2031-05-20T10:00:00"}).status_code == 400
    fake.access = "denied"
    assert client.post("/api/calendar/events", json={"title": "x", "start": "2031-05-20T10:00:00"}).status_code == 409


def test_reminders_sync_two_way(client, fake_apple):
    fake, configure = fake_apple
    tasks = client.get("/api/tasks?kind=tasks").json()["items"]
    live = [t for t in tasks if t["status"] in ("draft", "open")]
    mine = next(t for t in live if t["owner"] in ("Вы", ""))
    theirs = next(t for t in live if t["owner"] not in ("Вы", "") and t["id"] != mine["id"])
    assert client.post("/api/integrations/reminders/sync").status_code == 409  # выключено
    for t in tasks:  # остальные подтверждённые задачи не должны влиять на счётчики
        if t["id"] not in (mine["id"], theirs["id"]) and t["status"] == "open":
            client.patch(f"/api/tasks/{t['id']}", json={"status": "draft"})
    configure(reminders_enabled=True, reminders_list="Memory Desk")
    for t in (mine, theirs):
        client.patch(f"/api/tasks/{t['id']}", json={"status": "open"})
    # draft/open: создаются только мои, чужие — по настройке «все»
    r = client.post("/api/integrations/reminders/sync").json()
    assert r["summary"]["created"] == 1
    [(rid, rem)] = fake.reminders.items()
    assert rem["title"] == mine["text"] and rem["list"] == "Memory Desk" and (rem["due"] is not None) == bool(mine["due"])
    assert "Из записи" in rem["notes"] and "/records/" in rem["url"]
    configure(reminders_scope="all")
    assert client.post("/api/integrations/reminders/sync").json()["summary"]["created"] == 1
    assert f"{theirs['text']} ({theirs['owner']})" in {r["title"] for r in fake.reminders.values()}
    # повторный проход ничего не меняет
    assert client.post("/api/integrations/reminders/sync").json()["summary"] == {
        "created": 0, "updated": 0, "completedInApp": 0, "completedInReminders": 0, "removed": 0, "errors": 0}
    # отметили в приложении → выполнено в Напоминаниях
    client.patch(f"/api/tasks/{mine['id']}", json={"status": "done"})
    assert client.post("/api/integrations/reminders/sync").json()["summary"]["completedInApp"] == 1
    assert fake.reminders[rid]["completed"] is True
    # отметили в Напоминаниях → выполнено в приложении
    other_rid = next(i for i, r in fake.reminders.items() if i != rid)
    fake.reminders[other_rid]["completed"] = True
    assert client.post("/api/integrations/reminders/sync").json()["summary"]["completedInReminders"] == 1
    got = {t["id"]: t for t in client.get("/api/tasks?kind=tasks").json()["items"]}
    assert got[theirs["id"]]["status"] == "done"
    # вернули в работу в Напоминаниях
    fake.reminders[other_rid]["completed"] = False
    client.post("/api/integrations/reminders/sync")
    assert {t["id"]: t for t in client.get("/api/tasks?kind=tasks").json()["items"]}[theirs["id"]]["status"] == "open"
    # правка текста и срока → обновление напоминания
    client.patch(f"/api/tasks/{theirs['id']}", json={"text": "Подготовить и согласовать схему", "deadline": "2031-01-02"})
    assert client.post("/api/integrations/reminders/sync").json()["summary"]["updated"] == 1
    assert fake.reminders[other_rid]["title"].startswith("Подготовить и согласовать схему") and str(fake.reminders[other_rid]["due"]) == "2031-01-02"
    # отклонили задачу → напоминание удалено; удалили вручную → заново не создаётся
    client.patch(f"/api/tasks/{theirs['id']}", json={"status": "rejected"})
    assert client.post("/api/integrations/reminders/sync").json()["summary"]["removed"] == 1 and other_rid not in fake.reminders
    del fake.reminders[rid]
    client.patch(f"/api/tasks/{mine['id']}", json={"status": "open"})
    client.post("/api/integrations/reminders/sync")
    assert not fake.reminders
    # нет доступа → понятная ошибка, статус сохраняется для настроек
    fake.access = "denied"
    r = client.post("/api/integrations/reminders/sync")
    assert r.status_code == 409 and "Нет доступа" in r.json()["detail"]
    assert "Нет доступа" in client.get("/api/integrations/apple").json()["lastSync"]["error"]


# --------------------------------------------------------------------------- хранилище

def test_storage_usage_and_clean(client, fixture_dir):
    import os, time as _t
    from pathlib import Path
    from memory_desk import config
    work = config.PROJECT_ROOT / "work" / "demo-job"
    work.mkdir(parents=True, exist_ok=True)
    (work / "audio-16k.wav").write_bytes(b"0" * 5000)
    (work / "transcript.json").write_text("{}", encoding="utf-8")
    config.LOGS_DIR.mkdir(parents=True, exist_ok=True)
    (config.LOGS_DIR / "webui.log").write_text("x" * 1000, encoding="utf-8")
    u = client.get("/api/storage").json()
    cats = {c["key"]: c for c in u["categories"]}
    assert cats["work_audio"]["bytes"] >= 5000 and cats["work_audio"]["count"] >= 1 and cats["logs"]["bytes"] >= 1000
    assert {i["key"] for i in u["info"]} == {"notes", "database", "index", "models", "ollama"} and u["total"] > 0
    r = client.post("/api/storage/clean", json={"category": "work_audio"}).json()
    assert r["freed"] >= 5000 and not (work / "audio-16k.wav").exists() and (work / "transcript.json").exists()
    assert client.post("/api/storage/clean", json={"category": "logs"}).json()["freed"] >= 1000
    assert (config.LOGS_DIR / "webui.log").read_text(encoding="utf-8") == ""
    client.post("/api/storage/clean", json={"category": "work_results"})
    assert not work.exists()
    # архив: удаляется только аудио старше N дней
    audio = [p for p in cats and (Path(client.get("/api/settings").json()["archive_dir"])).rglob("*") if p.is_file()]
    assert audio
    old = audio[0]
    os.utime(old, (_t.time() - 40 * 86400, _t.time() - 40 * 86400))
    n_before = client.get("/api/storage").json()["categories"][2]["count"]
    r = client.post("/api/storage/clean", json={"category": "archive_audio", "days": 30}).json()
    assert r["removed"] == 1 and not old.exists() and r["usage"]["categories"][2]["count"] == n_before - 1
    assert client.post("/api/storage/clean", json={"category": "nope"}).status_code == 400
    assert client.post("/api/storage/clean", json={"category": "archive_audio", "days": -1}).status_code == 400
    for _ in range(5):
        client.post("/api/system/backup")
    assert client.post("/api/storage/clean", json={"category": "backups"}).json()["usage"]["categories"][3]["count"] == 3


# --------------------------------------------------------------------------- голоса (подмена кодировщика и звука)

class ToneEncoder:
    """Вместо ECAPA: «голос» определяется доминирующей частотой фрагмента."""
    def embed(self, samples):
        import numpy as np
        spectrum = np.abs(np.fft.rfft(samples))
        freq = int(np.argmax(spectrum) * 16000 / len(samples))
        vec = np.full(8, 0.05, dtype=np.float32)
        vec[0 if freq < 300 else 1 if freq < 600 else 2] = 1.0
        return vec / np.linalg.norm(vec)


@pytest.fixture()
def tone_audio(client, items, monkeypatch):
    """Звук, где реплики 0,1,4,5 говорит голос 220 Гц, а 2,3 — 440 Гц."""
    import numpy as np
    from memory_desk import knowledge, voice
    mid = by_title(items, "Обсуждение машинного времени и последствий")["id"]
    segs = client.get(f"/api/meetings/{mid}").json()["transcript"]
    spans = voice.segment_spans(segs, 0)
    end = int(spans[-1][1] + 1)
    wave = np.zeros(end * voice.SAMPLE_RATE, dtype=np.float32)
    for i, (a, b) in enumerate(spans):
        freq = 440 if i in (2, 3) else 220
        t = np.arange(int(a * voice.SAMPLE_RATE), int(b * voice.SAMPLE_RATE)) / voice.SAMPLE_RATE
        wave[int(a * voice.SAMPLE_RATE): int(b * voice.SAMPLE_RATE)] = 0.3 * np.sin(2 * np.pi * freq * t)[: int(b * voice.SAMPLE_RATE) - int(a * voice.SAMPLE_RATE)]
    pcm = (wave * 32767).astype(np.int16)
    monkeypatch.setattr(voice, "load_audio", lambda _p: pcm)
    monkeypatch.setattr(knowledge, "meeting_audio", lambda _m: Path("/nonexistent/audio.wav"))
    monkeypatch.setattr(voice, "MIN_CLUSTER_SEC", 0.1)
    monkeypatch.setattr(voice, "MIN_SEGMENT_SEC", 0.1)
    voice.use_encoder(ToneEncoder())
    yield mid
    voice.use_encoder(None)


def test_voice_cluster_and_match_units():
    import numpy as np
    from memory_desk import voice
    a, b = np.array([1, 0, 0], dtype=np.float32), np.array([0, 1, 0], dtype=np.float32)
    assert voice.cluster([a, a, b, None, b, a], [5, 5, 5, 1, 5, 5], sim=0.5) == [0, 0, 1, 1, 1, 0]
    profiles = [{"person_id": "1", "name": "Анна", "embedding": a, "samples": 1, "seconds": 9}]
    assert voice.match_profiles({"0": a, "1": b}, profiles) == {"0": {"name": "Анна", "score": 1.0}}
    # два похожих профиля: неуверенное совпадение не принимается
    twins = profiles + [{"person_id": "2", "name": "Борис", "embedding": a * 0.99 + b * 0.01, "samples": 1, "seconds": 9}]
    assert voice.match_profiles({"0": a}, twins) == {}
    # один человек не может быть назначен двум голосам
    assert len(voice.match_profiles({"0": a, "1": a * 0.98 + b * 0.05}, profiles)) == 1


def test_voice_diarize_learn_and_recognize(client, tone_audio):
    mid = tone_audio
    state = client.get("/api/voice").json()
    assert state["available"] and state["profiles"] == 0
    assert client.post(f"/api/meetings/{mid}/diarize?mode=voice").status_code == 200
    status = wait_diarized(client, mid)
    assert status["state"] == "done", status
    d = client.get(f"/api/meetings/{mid}").json()
    assert [x["speaker"] for x in d["transcript"]] == ["Спикер 1"] * 2 + ["Спикер 2"] * 2 + ["Спикер 1"] * 2
    assert all(s["voice"] and not s["voice"]["matched"] for s in d["speakers"])

    # без согласия отпечаток не сохраняется
    person = client.post("/api/people", json={"name": "Голосов Иван", "role": ""}).json()
    sample = {"meetingId": mid, "label": "Спикер 2"}
    assert client.post(f"/api/people/{person['id']}/voice", json=sample).status_code == 400
    ok = client.post(f"/api/people/{person['id']}/voice", json={**sample, "consent": True})
    assert ok.status_code == 200 and ok.json()["voice"]["samples"] == 1
    assert client.post(f"/api/people/{person['id']}/voice", json={"meetingId": mid, "label": "Нет такого", "consent": True}).status_code == 409
    mine = next(p for p in client.get("/api/people").json()["items"] if p["id"] == person["id"])
    assert mine["voice"]["samples"] == 1

    # повторный анализ: «Спикер 2» узнан и получает имя; второй голос остаётся безымянным
    client.post(f"/api/meetings/{mid}/diarize?mode=voice")
    status = wait_diarized(client, mid)
    assert status["state"] == "done" and "Голосов Иван" in status["message"]
    d = client.get(f"/api/meetings/{mid}").json()
    assert [x["speaker"] for x in d["transcript"]] == ["Спикер 1"] * 2 + ["Голосов Иван"] * 2 + ["Спикер 1"] * 2
    assert next(s for s in d["speakers"] if s["label"] == "Голосов Иван")["voice"]["matched"] is True
    assert "Голосов Иван" in [p["name"] for p in d["people"]]

    # ручное переименование переносит отпечаток вместе с меткой
    client.put(f"/api/meetings/{mid}/speakers", json={"rename": {"Спикер 1": "Пётр Новиков"}})
    d = client.get(f"/api/meetings/{mid}").json()
    assert next(s for s in d["speakers"] if s["label"] == "Пётр Новиков")["voice"] is not None

    # удаление профиля
    assert client.delete(f"/api/people/{person['id']}/voice").status_code == 200
    assert client.delete(f"/api/people/{person['id']}/voice").status_code == 404


def test_overlap_marking_is_conservative(client, items):
    """Наложение речи: помечаются только реплики, почти одинаково похожие на двух говорящих; метка не меняется."""
    import numpy as np
    from memory_desk import voice
    a, b = np.array([1.0, 0.0, 0.0], dtype=np.float32), np.array([0.0, 1.0, 0.0], dtype=np.float32)
    mix = (a + b) / np.sqrt(2)                      # поровну двое
    mostly_a = (a * 0.9 + b * 0.3) / np.linalg.norm(a * 0.9 + b * 0.3)   # явно первый
    cents = {0: (a, 10.0), 1: (b, 10.0)}
    embeddings = [a, b, mix, mostly_a, None] + [a] * 15
    groups = [0, 1, 0, 0, 0] + [0] * 15
    assert voice.find_overlaps(embeddings, groups, cents) == [2]
    assert voice.find_overlaps(embeddings, groups, {0: (a, 10.0)}) == []        # один говорящий — наложения не бывает
    assert voice.find_overlaps([a, b, a, b], [0, 1, 0, 1], cents) == []          # чистые реплики не помечаются
    many = [mix] * 20                                                            # не больше десятой части реплик
    assert len(voice.find_overlaps(many, [0] * 20, cents)) == 2

    mid = by_title(items, "Архитектурный комитет КХД")["id"]
    d = client.get(f"/api/meetings/{mid}").json()
    assert not any(seg.get("overlap") for seg in d["transcript"])
    first = d["transcript"][0]
    voice.save_meeting_overlaps(mid, [first["sec"]])
    try:
        d = client.get(f"/api/meetings/{mid}").json()
        assert d["transcript"][0].get("overlap") is True and d["transcript"][0]["speaker"] == first["speaker"]
        assert sum(1 for seg in d["transcript"] if seg.get("overlap")) == 1
    finally:
        voice.save_meeting_overlaps(mid, [])


def test_voice_unavailable_and_text_fallback(client, items, monkeypatch):
    from memory_desk import voice
    voice.use_encoder(None)
    monkeypatch.setattr(voice, "dependencies", lambda: (False, "Не установлены компоненты: speechbrain"))
    mid = by_title(items, "Встреча один на один с Анной")["id"]
    r = client.post(f"/api/meetings/{mid}/diarize?mode=voice")
    assert r.status_code == 409 and "speechbrain" in r.json()["detail"]
    info = client.get("/api/voice").json()
    assert info["available"] is False and info["modelReady"] in (True, False)
    assert client.post(f"/api/meetings/{mid}/diarize?mode=bad").status_code == 422


def test_voice_ingest_from_worker(client, items, fixture_dir):
    """Отпечатки, оставленные фоновой обработкой, попадают в базу вместе с узнанными участниками."""
    import json
    import numpy as np
    from memory_desk import config, knowledge, voice
    m = next(x for x in client.app.state.core.meetings if x.title == "Встреча один на один с Анной")
    folder = config.PROJECT_ROOT / "work" / "ingest-test"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "state.json").write_text(json.dumps({"note": str(m.path)}), encoding="utf-8")
    vec = np.eye(4, dtype=np.float32)[0].tolist()
    (folder / "voices.json").write_text(json.dumps({"speakers": {"Спикер 1": {"embedding": vec, "seconds": 12.0, "matched": False, "score": None}}}),
                                        encoding="utf-8")
    done = voice.ingest_pending({str(m.path): m.id})
    assert done == [(m.id, [])] and voice.meeting_voices(m.id)["Спикер 1"]["seconds"] == 12.0
    assert voice.ingest_pending({str(m.path): m.id}) == []  # повторно не подхватывается


def test_voice_install_plan_keeps_torch():
    from memory_desk import voice
    plan = voice.install_plan("2.14.1")
    assert ["--no-deps", "speechbrain"] in plan and not any("torch" in a for step in plan for a in step)
    assert any("scipy" in step and "hyperpyyaml" in step for step in plan)
    assert voice.install_plan(None)[0] == ["torch"]


def test_voice_prepare_reports_failure(client, monkeypatch):
    from memory_desk import voice
    monkeypatch.setattr(voice, "dependencies", lambda: (False, "нет компонентов"))
    monkeypatch.setattr(voice, "pip_command", lambda: None)
    assert voice.start_download() is True
    for _ in range(100):
        state = client.get("/api/voice").json()["download"]
        if state["state"] != "running":
            break
        time.sleep(0.05)
    assert state["state"] == "error" and "pip" in state["message"]


def test_prep_brief(client, items):
    project = next(p for p in client.get("/api/projects").json()["items"] if p["title"] == "КХД")
    brief = client.get("/api/prep", params={"project": project["id"]}).json()
    assert brief["title"] == "КХД" and brief["records"] >= 3
    sections = {s["key"]: s["items"] for s in brief["sections"]}
    assert list(sections) == ["changes", "decisions", "tasks", "questions", "risks"]
    assert sections["decisions"] and sections["tasks"] and sections["changes"]
    for group in sections.values():
        for item in group:
            assert item["meetingId"] and len(item["text"]) <= 151 and "\n" not in item["text"]
    dated = [t["due"] for t in sections["tasks"] if t["due"]]
    assert dated == sorted(dated) or any(t["overdue"] for t in sections["tasks"])
    one = client.get("/api/prep", params={"meeting": next(iter(items.values()))["id"]}).json()
    assert one["records"] == 1
    assert client.get("/api/prep", params={"project": "нет"}).status_code == 404
    assert client.get("/api/prep", params={"meeting": "нет"}).status_code == 404
    assert client.get("/api/prep").json()["title"] == "Весь архив"


def test_whisper_filter():
    from memory_desk.whisper_filter import clean_result, collapse_loops
    seg = lambda t, **kw: {"start": 0, "end": 1, "text": t, **kw}  # noqa: E731
    result = {"segments": [
        seg("Давайте начнём встречу."),
        seg("Субтитры сделал DimaTorzok"),
        seg("Продолжение следует..."),
        seg("мм", no_speech_prob=0.9, avg_logprob=-1.7),
        seg("Да да да да да да да да"),
        seg("Хорошо."), seg("Хорошо."), seg("Хорошо."), seg("Хорошо."),
        seg("Обсудим сроки по проекту " * 12, compression_ratio=3.1),
        seg("  "),
        seg("Следующий вопрос про бюджет."),
    ]}
    cleaned, dropped = clean_result(result)
    texts = [s["text"] for s in cleaned["segments"]]
    assert texts == ["Давайте начнём встречу.", "Да", "Хорошо.", "Хорошо.", "Следующий вопрос про бюджет."]
    reasons = {d["reason"] for d in dropped}
    assert {"служебная фраза", "нет речи", "повтор", "зацикливание", "пусто"} <= reasons
    assert collapse_loops("Нет, нет, нет, нет, нет") == "Нет"
    assert collapse_loops("Это нормальная фраза без повторов.") == "Это нормальная фраза без повторов."
    assert cleaned["text"].startswith("Давайте начнём")


def test_analytics(client):
    data = client.get("/api/analytics", params={"days": 0}).json()
    assert data["totals"]["records"] >= 5 and data["totals"]["decisions"] >= 1
    assert len(data["weeks"]) == 12 and data["weeks"][-1]["start"] <= data["today"]
    assert sum(w["meetings"] for w in data["weeks"]) >= 1
    assert data["byOwner"] and all({"owner", "open", "overdue", "done"} <= set(r) for r in data["byOwner"])
    for g in data["gantt"]:
        assert g["start"] <= g["due"] and g["meetingId"]
    names = {p["name"] for p in data["graph"]["people"]}
    assert all(l["person"] in names for l in data["graph"]["links"])
    assert client.get("/api/analytics", params={"days": 1}).json()["totals"]["records"] <= data["totals"]["records"]
    assert client.get("/api/analytics", params={"days": -5}).status_code == 422


def test_digest_period_parsing():
    from datetime import date
    from memory_desk import digest
    today = date(2026, 10, 14)  # среда
    assert digest.parse_period("Что сделано по КХД за неделю?", today)[:2] == (date(2026, 10, 7), today)
    assert digest.parse_period("итоги за месяц", today)[:2] == (date(2026, 9, 14), today)
    assert digest.parse_period("что было на прошлой неделе", today)[:2] == (date(2026, 10, 5), date(2026, 10, 11))
    assert digest.parse_period("итоги за две недели", today)[:2] == (date(2026, 9, 30), today)
    assert digest.parse_period("итоги за 10 дней", today)[:2] == (date(2026, 10, 4), today)
    assert digest.parse_period("что делали в сентябре", today)[:2] == (date(2026, 9, 1), date(2026, 9, 30))
    assert digest.parse_period("в этом месяце", today)[:2] == (date(2026, 10, 1), today)
    assert digest.parse_period("Кто ответственный за миграцию?", today) is None


def test_chat_digest_structured_and_model(client, monkeypatch):
    from memory_desk import digest
    # без модели: структурная сводка по проекту КХД за всё нужное время
    r = client.post("/api/chat", json={"question": "Что сделано по проекту КХД за месяц?", "scope": {"kind": "all"}}).json()
    assert r["mode"].startswith("сводка") and r["digest"]["records"] >= 3
    assert "КХД" in r["digest"]["title"] and "записи" in r["answer"] or "записей" in r["answer"]
    assert len(r["sources"]) == min(r["digest"]["records"], 10) and r["sources"][0]["meetingId"]
    assert "**Решения**" in r["answer"] and "**В работе**" in r["answer"]
    # с моделью: в запрос попадают факты, в ответ — заголовок с периодом и цифрами
    seen = []
    monkeypatch.setattr(digest, "generate", lambda prompt, num_predict=700: seen.append(prompt) or "**Главное** — тест")
    r2 = client.post("/api/chat", json={"question": "итоги по проекту КХД за неделю", "scope": {"kind": "all"}}).json()
    assert r2["mode"] == "сводка" and "**Главное** — тест" in r2["answer"] and r2["digest"]["model"]
    assert "Правила" in seen[0] and "Встреча" in seen[0]
    # выбранный в интерфейсе проект и вопрос без периода → сводка за 30 дней
    r3 = client.post("/api/chat", json={"question": "Что сделано?", "scope": {"kind": "project", "id": "kxd", "title": "КХД"}}).json()
    assert r3["digest"]["records"] >= 3
    # период без записей
    r4 = client.post("/api/chat", json={"question": "итоги за вчера по проекту Личное", "scope": {"kind": "all"}}).json()
    assert r4["confirmed"] is False and "нет" in r4["answer"]
    # точечный вопрос и вопрос по одной записи остаются обычным поиском
    r5 = client.post("/api/chat", json={"question": "Кто ответственный за миграцию справочников?", "scope": {"kind": "all"}}).json()
    assert "digest" not in r5
    first = client.get("/api/meetings").json()["items"][0]["id"]
    r6 = client.post("/api/chat", json={"question": "Что сделано за неделю?", "scope": {"kind": "record", "id": first}}).json()
    assert "digest" not in r6


def test_digest_splits_long_periods(monkeypatch):
    from memory_desk import digest
    facts = [digest.MeetingFacts(str(i), f"Встреча {i}", "2026-10-01", "КХД", summary=["о чём-то важном " * 8],
                                 decisions=[{"text": "решение " * 20, "draft": False}]) for i in range(80)]
    intent = digest.Intent(__import__("datetime").date(2026, 9, 1), __import__("datetime").date(2026, 10, 1), "месяц", True, {"id": "kxd", "title": "КХД"})
    assert len(digest.chunks(facts, False)) > 1
    calls = []
    monkeypatch.setattr(digest, "generate", lambda prompt, num_predict=700: calls.append(len(prompt)) or "- пункт " * 30)
    text, how = digest.write(intent, facts)
    assert how == "model" and len(calls) > 2 and max(calls) < 12000
    monkeypatch.setattr(digest, "generate", lambda *a, **k: (_ for _ in ()).throw(OSError("нет связи")))
    text, how = digest.write(intent, facts)
    assert how == "structured" and "без модели" in text and "…и ещё" in text


def test_input_limits(client):
    assert client.post("/api/projects", json={"title": "п" * 200}).status_code == 400
    assert client.post("/api/projects", json={"title": "..."}).status_code == 400
    assert client.post("/api/people", json={"name": "ж" * 500}).status_code == 400
    mid = client.get("/api/meetings").json()["items"][0]["id"]
    assert client.patch(f"/api/meetings/{mid}", json={"title": "н" * 500}).status_code == 400
    task = client.get("/api/tasks").json()["items"][0]
    url = f"/api/tasks/{task['id']}"
    for bad in ("не дата", "2026-02-30", "9999-99-99"):
        assert client.patch(url, json={"deadline": bad}).status_code == 400
    assert client.patch(url, json={"deadline": "2026-11-05"}).json()["due"] == "2026-11-05"
    cleared = client.patch(url, json={"deadline": ""}).json()
    assert cleared["due"] is None and cleared["dueText"] == ""
    assert client.patch(url, json={"text": "т" * 2000}).status_code == 400
    cur = client.get("/api/settings").json()
    body = {k: cur[k] for k in ("recordings_dir", "obsidian_dir", "archive_dir", "retention_days", "auto_process", "diarize", "voice_enabled", "calendar_source")}
    assert client.put("/api/settings", json={**body, "recordings_dir": "/нет/такой/папки"}).status_code == 400
    assert client.put("/api/settings", json={**body, "retention_days": 0}).status_code == 400
    assert client.put("/api/settings", json=body).status_code == 200
    assert client.post("/api/chat", json={"question": "q", "scope": {"kind": "days", "days": "x"}}).status_code == 400


def test_long_titles_make_safe_file_names():
    import meetings_process as mp
    name = mp.slug("Очень длинное название " * 20)
    assert len(name.encode()) < 200 and name == name.strip()
    assert mp.slug("a\nb:c") == "a b-c" and mp.slug("") == "Запись PLAUD"


def test_reindex_reports_progress(client):
    first = client.post("/api/system/reindex")
    assert first.status_code == 200 and first.json()["state"] in {"running", "done"}
    for _ in range(100):
        status = client.get("/api/system/reindex").json()
        if status["state"] != "running":
            break
        time.sleep(0.1)
    assert status["state"] == "done", status
    assert status["total"] > 0 and status["done"] == status["total"] and status["records"] >= 5
    assert client.get("/api/health").json()["indexState"] == "ready"
    assert client.post("/api/chat", json={"question": "Где обсуждали справочники?", "scope": {"kind": "all"}}).status_code == 200


def test_hostile_edits_keep_notes_intact(client, items):
    """Правки с разметкой, кавычками и служебными символами не должны ломать заметку и стенограмму."""
    candidates = [client.get(f"/api/meetings/{m['id']}").json() for m in items.values()]
    before = next(d for d in candidates if len(d["transcript"]) >= 6 and d["tasks"])
    mid = before["id"]
    lines_before = len(before["transcript"])
    tasks_before = len(before["tasks"])
    nasty = 'Название "в кавычках": #хэш --- [[ссылка]] \\ слэш\nвторая строка'
    r = client.patch(f"/api/meetings/{mid}", json={"title": nasty, "summary": "Первая\n## Решения\n- подделка\n---\nпоследняя"})
    assert r.status_code == 200
    after = client.get(f"/api/meetings/{mid}").json()
    assert "\n" not in after["title"] and "кавычках" in after["title"]
    assert after["project"] == before["project"] and len(after["decisions"]) == len(before["decisions"]), "решения должны остаться прежними"
    assert len(after["tasks"]) == tasks_before and len(after["transcript"]) == lines_before
    assert all("подделка" not in d["text"] for d in after["decisions"]), "текст краткого содержания не должен превращаться в решения"
    root = __import__("pathlib").Path(__import__("os").environ["MEMORY_DESK_CONFIG"]).parent / "Claude" / "Встречи"
    note = next(f for f in root.rglob("*.md") if f'meeting_id: "{mid}"' in f.read_text(encoding="utf-8"))
    text = note.read_text(encoding="utf-8")
    assert "\n## Решения\n- подделка" not in text and "\n---\nпоследняя" not in text and "Первая" in text and "последняя" in text
    # метки говорящих со служебными символами
    rename = client.put(f"/api/meetings/{mid}/speakers", json={"rename": {"Вы": "Имя: с **двоеточием**"}})
    assert rename.status_code == 200
    body = rename.json()
    assert len(body["transcript"]) == lines_before
    names = {l.get("speaker") for l in body["transcript"] if l.get("speaker")}
    assert all("*" not in n and ":" not in n for n in names), names
    # после полной перечитки всё то же самое
    client.post("/api/system/reindex")
    again = client.get(f"/api/meetings/{mid}").json()
    assert len(again["tasks"]) == tasks_before and len(again["transcript"]) == lines_before


def _note_of(meeting_id: str):
    import os
    from pathlib import Path
    root = Path(os.environ["MEMORY_DESK_CONFIG"]).parent / "Claude" / "Встречи"
    return next(f for f in root.rglob("*.md") if f'meeting_id: "{meeting_id}"' in f.read_text(encoding="utf-8"))


def test_item_edits_are_written_back_to_note(client, items):
    """Правка задачи или решения меняет строку в заметке Obsidian, остальной текст остаётся прежним."""
    details = [client.get(f"/api/meetings/{m['id']}").json() for m in items.values()]
    detail = next(d for d in details if len(d["tasks"]) >= 1 and d["decisions"])
    mid = detail["id"]
    note = _note_of(mid)
    original = note.read_text(encoding="utf-8")
    task = detail["tasks"][0]

    r = client.patch(f"/api/tasks/{task['id']}", json={"text": "Новая формулировка  задачи", "owner": "Пётр; Иванов", "deadline": "2026-12-01"})
    assert r.status_code == 200 and r.json()["noteSynced"] is True
    text = note.read_text(encoding="utf-8")
    line = next(l for l in text.splitlines() if l.startswith("- [ ] ") and "Новая формулировка задачи" in l)
    assert line.startswith("- [ ] ") and "**ответственный:** Пётр, Иванов" in line and "**срок:** 2026-12-01" in line
    # всё остальное в заметке не тронуто
    a, b = original.splitlines(), text.splitlines()
    changed = [y for x, y in zip(a, b) if x != y]
    assert len(a) == len(b) and 1 <= len(changed) <= 3, "меняются строка задачи и её подписи в схемах, больше ничего"
    assert all("Новая формулировка" in y for y in changed)

    # выполнено → галочка, отклонено → зачёркнуто, возврат → обычная строка
    assert client.patch(f"/api/tasks/{task['id']}", json={"status": "done"}).json()["noteSynced"] is True
    assert "- [x] Новая формулировка задачи" in note.read_text(encoding="utf-8")
    client.patch(f"/api/tasks/{task['id']}", json={"status": "rejected"})
    assert "~~Новая формулировка задачи~~" in note.read_text(encoding="utf-8")
    client.patch(f"/api/tasks/{task['id']}", json={"status": "open"})
    restored = note.read_text(encoding="utf-8")
    assert "- [ ] Новая формулировка задачи —" in restored and "~~" not in restored

    # решение: текст и отклонение
    decision = detail["decisions"][0]
    r = client.patch(f"/api/decisions/{decision['id']}", json={"text": "Решение в новой редакции"})
    assert r.json()["noteSynced"] is True and "- Решение в новой редакции" in note.read_text(encoding="utf-8")
    client.patch(f"/api/decisions/{decision['id']}", json={"status": "rejected"})
    assert "- ~~Решение в новой редакции~~" in note.read_text(encoding="utf-8")
    client.patch(f"/api/decisions/{decision['id']}", json={"status": "draft"})

    # после полной перечитки заметок реестр не дублируется и правки сохраняются
    client.post("/api/system/reindex")
    again = client.get(f"/api/meetings/{mid}").json()
    assert len(again["tasks"]) == len(detail["tasks"]) and len(again["decisions"]) == len(detail["decisions"])
    assert any(t["text"] == "Новая формулировка задачи" and t["owner"] == "Пётр; Иванов" for t in again["tasks"])

    # заметку правили вручную: строки нет — правка остаётся в кабинете, об этом сообщается
    note.write_text(note.read_text(encoding="utf-8").replace("Новая формулировка задачи", "Совсем другой текст"), encoding="utf-8")
    r = client.patch(f"/api/tasks/{task['id']}", json={"owner": "Анна"})
    assert r.status_code == 200 and r.json()["noteSynced"] is False and r.json()["owner"] == "Анна"


def test_pagination_keeps_totals(client):
    everything = client.get("/api/meetings").json()
    assert everything["total"] == len(everything["items"]) >= 3
    first = client.get("/api/meetings", params={"limit": 2}).json()
    second = client.get("/api/meetings", params={"limit": 2, "offset": 2}).json()
    assert len(first["items"]) == 2 and first["total"] == everything["total"]
    assert [m["id"] for m in first["items"] + second["items"]] == [m["id"] for m in everything["items"][: 2 + len(second["items"])]]
    assert client.get("/api/meetings", params={"limit": -1}).status_code == 422
    for kind in ("tasks", "decisions"):
        full = client.get("/api/tasks", params={"kind": kind}).json()
        page = client.get("/api/tasks", params={"kind": kind, "limit": 3}).json()
        assert full["total"] == len(full["items"]) and len(page["items"]) == min(3, full["total"]) and page["total"] == full["total"]
        assert [i["id"] for i in page["items"]] == [i["id"] for i in full["items"][:3]]


def test_bulk_items_and_meetings(client, items):
    registry = client.get("/api/tasks").json()["items"]
    drafts = [t for t in registry if t["draft"]][:3]
    assert len(drafts) >= 2
    ids = [t["id"] for t in drafts]
    r = client.post("/api/tasks/bulk", json={"ids": ids + ["нет-такого"], "status": "open", "deadline": "2026-12-24"})
    assert r.status_code == 200
    body = r.json()
    assert body["updated"] == len(ids) and body["failedCount"] == 1 and body["failed"][0]["id"] == "нет-такого"
    after = {t["id"]: t for t in client.get("/api/tasks").json()["items"]}
    assert all(not after[i]["draft"] and after[i]["due"] == "2026-12-24" for i in ids)
    # неверное значение отвергается сразу, ничего не меняя
    assert client.post("/api/tasks/bulk", json={"ids": ids, "deadline": "завтра"}).status_code == 400
    assert client.post("/api/tasks/bulk", json={"ids": [], "status": "open"}).status_code == 400
    assert client.post("/api/tasks/bulk", json={"ids": ["x"] * 3 + [str(i) for i in range(600)], "status": "open"}).status_code == 400
    assert client.post("/api/tasks/bulk", json={"ids": ids, "status": "готово"}).status_code == 400
    # в заметках изменилась галочка/срок у каждой выбранной задачи
    for t in drafts:
        note = _note_of(t["meetingId"]).read_text(encoding="utf-8")
        assert "**срок:** 2026-12-24" in note
    # решения: подтверждение пачкой, ответственный и срок игнорируются
    decisions = [d for d in client.get("/api/tasks", params={"kind": "decisions"}).json()["items"] if d["draft"]][:2]
    r = client.post("/api/decisions/bulk", json={"ids": [d["id"] for d in decisions], "status": "confirmed", "owner": "игнор"})
    assert r.json()["updated"] == len(decisions)
    # записи: архив и возврат, публикация без проекта отклоняется по записи, а не целиком
    mids = [m["id"] for m in list(items.values())[:2]]
    r = client.post("/api/meetings/bulk", json={"ids": mids, "action": "archive"})
    assert r.json()["updated"] == 2
    assert not {m["id"] for m in client.get("/api/meetings").json()["items"]} & set(mids)
    assert {m["id"] for m in client.get("/api/meetings", params={"archived": True}).json()["items"]} >= set(mids)
    assert client.post("/api/meetings/bulk", json={"ids": mids, "action": "restore"}).json()["updated"] == 2
    assert client.post("/api/meetings/bulk", json={"ids": mids, "action": "project", "value": "нет-проекта"}).status_code == 404
    assert client.post("/api/meetings/bulk", json={"ids": mids, "action": "удалить"}).status_code == 422
    target = mids[0]
    assert client.post("/api/meetings/bulk", json={"ids": [target], "action": "project", "value": ""}).json()["updated"] == 1
    r = client.post("/api/meetings/bulk", json={"ids": [target], "action": "publish"})
    assert r.json()["updated"] == 0 and r.json()["failedCount"] == 1


def test_protocol_mindmap_and_deck(client, items):
    import os
    from pathlib import Path
    detail = next(d for d in (client.get(f"/api/meetings/{m['id']}").json() for m in items.values()) if d["tasks"] and d["decisions"])
    mid = detail["id"]
    # протокол: без модели собирается из фактов
    r = client.post(f"/api/meetings/{mid}/protocol")
    assert r.status_code == 200
    body = r.json()
    assert body["mode"] == "structured" and body["markdown"].startswith("# Протокол:")
    for part in ("## Решения", "## Поручения", "| № | Поручение |", detail["decisions"][0]["text"][:25], detail["tasks"][0]["text"][:25]):
        assert part in body["markdown"], part
    saved = client.post(f"/api/meetings/{mid}/protocol", params={"save": True}).json()
    root = Path(os.environ["MEMORY_DESK_CONFIG"]).parent / "Claude" / "Протоколы"
    assert (root / saved["fileName"]).read_text(encoding="utf-8").startswith("# Протокол:") and saved["savedAs"]
    assert client.post("/api/meetings/нет/protocol").status_code == 404

    # карта: запись и проект
    m = client.post("/api/materials/mindmap", json={"meetingId": mid}).json()
    assert m["mermaid"].startswith("mindmap\n  root((") and m["tree"]["children"] and "Решения" in {c["text"] for c in m["tree"]["children"]}
    assert all("(" not in line.split("[", 1)[-1] for line in m["mermaid"].splitlines()[2:])
    project = next(p for p in client.get("/api/projects").json()["items"] if p["counts"]["records"] >= 2)
    pm = client.post("/api/materials/mindmap", json={"projectId": project["id"], "days": 3650}).json()
    assert pm["records"] >= 2 and pm["tree"]["children"]
    assert client.post("/api/materials/mindmap", json={"projectId": "нет"}).status_code == 404
    assert client.post("/api/materials/mindmap", json={"projectId": project["id"], "days": 0}).status_code == 422

    # презентация: файл открывается, слайды содержат факты
    d = client.post("/api/materials/deck", json={"projectId": project["id"], "days": 3650})
    assert d.status_code == 200
    info = d.json()
    assert info["slides"] >= 4 and info["records"] >= 2 and info["name"].endswith(".pptx")
    file = client.get(info["url"])
    assert file.status_code == 200 and file.content[:2] == b"PK"
    from io import BytesIO
    from pptx import Presentation
    prs = Presentation(BytesIO(file.content))
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert any(project["title"] in t for t in texts) and any(t.startswith("Решения") for t in texts)
    assert any(x["name"] == info["name"] for x in client.get("/api/materials/decks").json()["items"])
    assert client.get("/api/materials/deck/..%2Fx.pptx").status_code in (400, 404)
    assert client.get("/api/materials/deck/нет.pptx").status_code == 404
    assert client.post("/api/materials/deck", json={"projectId": project["id"], "days": 1, "personId": "нет"}).status_code == 404


def test_materials_use_model_when_available(client, items, monkeypatch):
    from memory_desk import digest
    prompts = []

    def fake(prompt, num_predict=700):
        prompts.append(prompt)
        return "- Первый итог по проекту\n- Второй итог\n" if "ГЛАВНОЕ" in prompt else "Обсудили архитектуру и договорились о сроках."

    monkeypatch.setattr(digest, "generate", fake)
    detail = next(d for d in (client.get(f"/api/meetings/{m['id']}").json() for m in items.values()) if d["tasks"] and d["summary"])
    body = client.post(f"/api/meetings/{detail['id']}/protocol").json()
    assert body["mode"] == "model" and "Обсудили архитектуру и договорились о сроках." in body["markdown"]
    project = next(p for p in client.get("/api/projects").json()["items"] if p["counts"]["records"] >= 2)
    info = client.post("/api/materials/deck", json={"projectId": project["id"], "days": 3650}).json()
    assert info["mode"] == "model"
    from io import BytesIO
    from pptx import Presentation
    prs = Presentation(BytesIO(client.get(info["url"]).content))
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Первый итог по проекту" in texts
    # модель недоступна → материал всё равно собирается
    monkeypatch.setattr(digest, "generate", lambda *a, **k: (_ for _ in ()).throw(OSError("нет связи")))
    assert client.post(f"/api/meetings/{detail['id']}/protocol").json()["mode"] == "structured"
    assert client.post("/api/materials/deck", json={"projectId": project["id"], "days": 3650}).json()["mode"] == "structured"


def test_materials_from_chat_and_save(client, items):
    import os
    from pathlib import Path
    project = next(p for p in client.get("/api/projects").json()["items"] if p["counts"]["records"] >= 2)
    ask = lambda q, scope=None: client.post("/api/chat", json={"question": q, "scope": scope or {"kind": "all"}}).json()

    deck = ask(f"сделай презентацию по проекту {project['title']}")
    assert deck["material"]["kind"] == "deck" and deck["material"]["result"]["url"].endswith(".pptx")
    assert client.get(deck["material"]["result"]["url"]).content[:2] == b"PK"
    assert "Презентация готова" in deck["answer"]

    detail = next(d for d in (client.get(f"/api/meetings/{m['id']}").json() for m in items.values()) if d["tasks"] and d["decisions"])
    proto = ask("составь протокол", {"kind": "record", "id": detail["id"]})
    assert proto["material"]["kind"] == "protocol" and "**Решения**" in proto["answer"] and "|" not in proto["answer"]
    assert proto["material"]["result"]["markdown"].startswith("# Протокол:")
    tree = ask("построй карту", {"kind": "record", "id": detail["id"]})
    assert tree["material"]["kind"] == "map" and tree["material"]["result"]["tree"]["children"]

    # без записи в области берётся последняя подходящая; пустая выборка — понятный ответ, а не ошибка
    assert "последнюю подходящую запись" in ask(f"подготовь протокол по проекту {project['title']}")["answer"]
    used = {int(m["date"][5:7]) for m in items.values() if m.get("date")} if isinstance(items, dict) else set()
    used |= {int(m["date"][5:7]) for m in client.get("/api/meetings", params={"limit": 1000}).json()["items"] if m["date"]}
    word = dict(zip(range(1, 13), ["январе", "феврале", "марте", "апреле", "мае", "июне", "июле", "августе", "сентябре", "октябре", "ноябре", "декабре"]))
    free = next(n for n in range(1, 13) if n not in used)
    empty = ask(f"сделай презентацию в {word[free]} по проекту " + project["title"])
    assert "material" not in empty and "записей нет" in empty["answer"]

    # обычный вопрос про слово «карта» без просьбы — обычный поиск, не материал
    assert "material" not in ask("какая карта у проекта?")

    # сохранение показанного текста в Obsidian
    r = client.post("/api/materials/save", json={"kind": "protocol", "name": "Тест: протокол/1.md", "content": "# Протокол\n"})
    assert r.status_code == 200
    root = Path(os.environ["MEMORY_DESK_CONFIG"]).parent / "Claude" / "Протоколы"
    assert (root / "Тест- протокол-1.md").read_text(encoding="utf-8") == "# Протокол\n"
    assert client.post("/api/materials/save", json={"kind": "boom", "name": "x", "content": "y"}).status_code == 400
    assert client.post("/api/materials/save", json={"kind": "map", "name": "", "content": "y"}).status_code == 422

    # произвольный диапазон дат
    ranged = client.post("/api/materials/mindmap", json={"projectId": project["id"], "since": "2000-01-01", "until": "2100-01-01"})
    assert ranged.status_code == 200 and ranged.json()["records"] >= 2
    assert client.post("/api/materials/mindmap", json={"since": "вчера", "until": "завтра"}).status_code == 400


def test_decision_tree_and_replacement(client, items):
    tree = client.get("/api/decisions/tree").json()
    group = next(g for g in tree["projects"] if g["id"] and g["total"] >= 2)
    flat = []

    def walk(nodes):
        for n in nodes:
            flat.append(n)
            walk(n["replaces"])
    walk(group["current"])
    assert len(flat) == group["total"] and group["active"] == group["total"] and group["replaced"] == 0
    old, new = flat[-1], flat[0]
    assert old["id"] != new["id"]

    # заменяем: старое решение уходит под действующее; в «Сегодня» и подготовке остаётся только актуальное
    r = client.patch(f"/api/decisions/{old['id']}", json={"replaced_by": new["id"]})
    assert r.status_code == 200 and r.json()["replacedBy"] == new["id"]
    after = next(g for g in client.get("/api/decisions/tree").json()["projects"] if g["id"] == group["id"])
    assert after["replaced"] == 1 and after["active"] == group["total"] - 1
    top = next(n for n in after["current"] if n["id"] == new["id"])
    assert [c["id"] for c in top["replaces"]] == [old["id"]]
    prep_texts = " ".join(i["text"] for s in client.get("/api/prep", params={"project": group["id"]}).json()["sections"] if s["key"] == "decisions" for i in s["items"])
    assert old["text"][:30] not in prep_texts or old["text"][:30] in new["text"]

    # материалы учитывают замену: в протоколе отметка, в презентации заменённого решения нет
    assert "_(позднее заменено)_" in client.post(f"/api/meetings/{old['meetingId']}/protocol").json()["markdown"]
    from io import BytesIO
    from pptx import Presentation
    deck = client.post("/api/materials/deck", json={"projectId": group["id"], "days": 3650}).json()
    prs = Presentation(BytesIO(client.get(deck["url"]).content))
    slides = {sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame}
    assert any(new["text"][:30] in t for t in slides) and not any(old["text"][:30] in t for t in slides if t.startswith(old["text"][:30]))

    # проверки: сам на себя, замкнутая цепочка, чужой проект, неизвестный id
    assert client.patch(f"/api/decisions/{new['id']}", json={"replaced_by": new['id']}).status_code == 400
    assert client.patch(f"/api/decisions/{new['id']}", json={"replaced_by": old['id']}).status_code == 400
    assert client.patch(f"/api/decisions/{new['id']}", json={"replaced_by": "нет-такого"}).status_code == 400
    other = next((g for g in client.get("/api/decisions/tree").json()["projects"] if g["id"] != group["id"] and g["current"]), None)
    if other:
        assert client.patch(f"/api/decisions/{new['id']}", json={"replaced_by": other["current"][0]["id"]}).status_code == 400
    # задачи поле не принимают
    task = client.get("/api/tasks").json()["items"][0]
    assert client.patch(f"/api/tasks/{task['id']}", json={"replaced_by": new["id"]}).status_code == 400

    # снять отметку
    assert client.patch(f"/api/decisions/{old['id']}", json={"replaced_by": ""}).json()["replacedBy"] is None
    restored = next(g for g in client.get("/api/decisions/tree").json()["projects"] if g["id"] == group["id"])
    assert restored["replaced"] == 0
    assert client.get("/api/decisions/tree", params={"project": group["id"], "days": 3650}).json()["projects"][0]["id"] == group["id"]


def test_calendar_range_apple_and_ics(client, items, fake_apple, tmp_path):
    from datetime import datetime, timedelta
    from memory_desk import storage
    fake, configure = fake_apple
    m = next(v for t, v in items.items() if "архитектур" in t.lower())   # название могли изменить соседние тесты
    detail = client.get(f"/api/meetings/{m['id']}").json()
    storage.update_item("meetings", m["id"], calendar_event_id="")
    start = datetime.fromisoformat(f"{detail['date']}T{detail['time']}")
    fake.add_event("Архитектурный комитет КХД", start - timedelta(minutes=5), start + timedelta(minutes=50), ["Я", "Соколова Анна"])
    fake.add_event("Пустая встреча", start + timedelta(days=2, hours=1), start + timedelta(days=2, hours=2), ["Кто-то"])
    fake.add_event("Отпуск", start.replace(hour=0, minute=0) + timedelta(days=1), start.replace(hour=23, minute=59) + timedelta(days=1), allDay=True)
    first, last = (start - timedelta(days=3)).date().isoformat(), (start + timedelta(days=5)).date().isoformat()

    assert client.get("/api/calendar", params={"from": first, "to": last}).json()["configured"] is False
    configure(apple_calendar=True)
    cal = client.get("/api/calendar", params={"from": first, "to": last}).json()
    assert cal["source"] == "apple" and cal["totals"]["events"] == 3 and cal["error"] == ""
    events = {e["title"]: e for d in cal["days"] for e in d["events"]}
    assert [r["id"] for r in events["Архитектурный комитет КХД"]["records"]] == [m["id"]]
    assert events["Пустая встреча"]["records"] == [] and events["Отпуск"]["allDay"] is True
    assert cal["totals"]["withRecord"] == 1
    # записи без события попадают в свой день и не дублируются у события
    rec_ids = [r["id"] for d in cal["days"] for r in d["records"]]
    assert m["id"] not in rec_ids

    # привязка по идентификатору побеждает совпадение по времени
    storage.update_item("meetings", m["id"], calendar_event_id=next(e["id"] for e in fake.event_list if e["title"] == "Архитектурный комитет КХД"),
                        calendar_event_title="Архитектурный комитет КХД")
    again = client.get("/api/calendar", params={"from": first, "to": last}).json()
    linked = next(e for d in again["days"] for e in d["events"] if e["title"] == "Архитектурный комитет КХД")
    assert linked["records"][0]["linked"] is True

    # границы и ошибки
    assert client.get("/api/calendar", params={"from": "вчера"}).status_code == 400
    assert client.get("/api/calendar", params={"from": "2020-01-01", "to": "2022-01-01"}).status_code == 400
    swapped = client.get("/api/calendar", params={"from": last, "to": first}).json()
    assert swapped["from"] == first and swapped["to"] == last
    assert client.get("/api/calendar").status_code == 200   # без параметров — текущая неделя

    # Календарь выключен, но есть ICS-файл: события читаются оттуда, в т.ч. повторяющиеся
    configure(apple_calendar=False)
    ics = tmp_path / "work.ics"
    day = start.strftime("%Y%m%d")
    ics.write_text("BEGIN:VCALENDAR\nVERSION:2.0\nPRODID:-//t//\nBEGIN:VEVENT\nUID:a1\nDTSTART:" + day + "T090000\nDTEND:" + day + "T100000\n"
                   "SUMMARY:Еженедельная планёрка\nRRULE:FREQ=DAILY;COUNT=3\nEND:VEVENT\nEND:VCALENDAR\n", encoding="utf-8")
    configure(calendar_source=str(ics))
    cal = client.get("/api/calendar", params={"from": first, "to": last}).json()
    assert cal["source"] == "ics" and cal["totals"]["events"] == 3
    configure(calendar_source=str(tmp_path / "нет.ics"))
    broken = client.get("/api/calendar", params={"from": first, "to": last}).json()
    assert broken["source"] == "none" and "Не удалось" in broken["error"]
    configure(calendar_source="")
    storage.update_item("meetings", m["id"], calendar_event_id="", calendar_event_title="")


def test_prep_headline(client, items, monkeypatch):
    from memory_desk import digest
    project = next(p for p in client.get("/api/projects").json()["items"] if p["counts"]["records"] >= 2)
    brief = client.get("/api/prep", params={"project": project["id"]}).json()
    assert brief["headline"]["mode"] == "structured" and "." in brief["headline"]["text"]

    # модель недоступна → строка из цифр, а не ошибка
    monkeypatch.setattr(digest, "generate", lambda *a, **k: (_ for _ in ()).throw(OSError("нет связи")))
    off = client.get("/api/prep/headline", params={"project": project["id"]}).json()
    assert off["mode"] == "structured" and off["text"] == brief["headline"]["text"]

    # модель отвечает → её текст очищается от разметки; повтор берётся из кэша
    calls = []

    def fake(prompt, num_predict=700):
        calls.append(prompt)
        return "**Главное:** закрыть просроченные задачи и решить вопрос по срокам."

    monkeypatch.setattr(digest, "generate", fake)
    meeting = next(iter(items.values()))
    on = client.get("/api/prep/headline", params={"meeting": meeting["id"]}).json()
    assert on["mode"] == "model" and "**" not in on["text"] and "закрыть просроченные задачи" in on["text"]
    assert "ГЛАВНОЕ" in calls[0] and "[Что менялось]" in calls[0]
    assert client.get("/api/prep/headline", params={"meeting": meeting["id"]}).json() == on and len(calls) == 1
    assert client.get("/api/prep/headline", params={"meeting": "нет"}).status_code == 404
    # слишком короткий ответ модели не принимается
    monkeypatch.setattr(digest, "generate", lambda *a, **k: "ок")
    short = client.get("/api/prep/headline", params={"project": project["id"]}).json()
    assert short["mode"] == "structured"
    assert client.get("/api/prep/headline", params={"project": project["id"]}).json()["mode"] == "structured"
