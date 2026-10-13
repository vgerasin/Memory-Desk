"""Этап 4.4: дайджест по проектам за неделю и месяц, текстом и презентацией."""
from __future__ import annotations

import re
import sys
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
KEYS = ("this_week", "last_week", "this_month", "last_month")


@pytest.fixture(autouse=True)
def _modules(fixture_dir, monkeypatch):
    sys.path.insert(0, str(ROOT))
    import memory_desk.artifacts, memory_desk.digest, memory_desk.materials, memory_desk.projectdigest, memory_desk.services, memory_desk.storage  # noqa: F401
    globals().update(artifacts=memory_desk.artifacts, digest=memory_desk.digest, materials=memory_desk.materials,
                     projectdigest=memory_desk.projectdigest, services=memory_desk.services, storage=memory_desk.storage)
    monkeypatch.setattr(digest, "generate", lambda *a, **k: (_ for _ in ()).throw(OSError("нет связи")))


def all_meetings(client):
    return client.get("/api/meetings", params={"limit": 1000}).json()["items"]


def rich_meeting(client):
    for m in all_meetings(client):
        d = client.get(f"/api/meetings/{m['id']}").json()
        if d["decisions"] and d["tasks"] and d.get("project"):
            return d
    raise AssertionError("нет записи с проектом, решениями и задачами")


def digest_for(client, **body):
    r = client.post("/api/materials/digest", json=body)
    assert r.status_code == 200, r.text
    return r.json()


def busy_period(client):
    """Период (ключ), за который в тестовых данных есть записи, и сам дайджест."""
    for key in KEYS:
        r = client.post("/api/materials/digest", json={"period": key})
        if r.status_code == 200:
            return key, r.json()
    raise AssertionError("в тестовых данных нет записей ни за один из периодов")


# --------------------------------------------------------------------------- границы периодов

def test_period_ranges():
    wed = date(2026, 10, 14)     # среда
    assert projectdigest.period_range("this_week", wed)[:2] == (date(2026, 10, 12), wed)
    since, until, label = projectdigest.period_range("last_week", wed)
    assert (since, until) == (date(2026, 10, 5), date(2026, 10, 11)) and label == "Неделя 5–11 октября 2026"
    assert projectdigest.period_range("this_month", wed)[:2] == (date(2026, 10, 1), wed)
    since, until, label = projectdigest.period_range("last_month", wed)
    assert (since, until, label) == (date(2026, 9, 1), date(2026, 9, 30), "Сентябрь 2026")
    # через границу года и месяца
    assert projectdigest.period_range("last_month", date(2027, 1, 5))[:2] == (date(2026, 12, 1), date(2026, 12, 31))
    assert projectdigest.period_range("last_week", date(2026, 11, 3))[2] == "Неделя 26 октября – 1 ноября 2026"
    with pytest.raises(ValueError):
        projectdigest.period_range("year", wed)


def test_period_from_text():
    f = projectdigest.period_from_text
    mon, fri = date(2026, 10, 12), date(2026, 10, 16)
    assert f("сделай недельный дайджест", mon) == "last_week" and f("сделай недельный дайджест", fri) == "this_week"
    assert f("дайджест за прошлую неделю", fri) == "last_week" and f("дайджест за эту неделю", mon) == "this_week"
    assert f("месячный дайджест", date(2026, 10, 20)) == "this_month" and f("месячный дайджест", date(2026, 10, 2)) == "last_month"
    assert f("дайджест за прошлый месяц", date(2026, 10, 20)) == "last_month"


# --------------------------------------------------------------------------- документ

def test_digest_structure_and_timecodes(client):
    key, out = busy_period(client)
    md = out["markdown"]
    assert md.startswith("# Дайджест: ") and out["period"]["key"] == key
    for head in ("**Период:**", "**Записей:**", "## Главное"):
        assert head in md
    assert artifacts.untimed(md) == []
    assert out["mode"] == "structured" and out["records"] >= 1 and out["count"] >= 1
    # каждый проект — свой раздел, в нём сводная строка с цифрами
    for p in out["projects"]:
        assert f"## {p['title']}" in md
    assert all(re.search(r"⏱ \d{2}:\d{2}:\d{2}", l) for l in md.splitlines() if l.startswith("- ") and "**" not in l)


def test_only_records_of_the_period_and_counts_match(client):
    key, out = busy_period(client)
    since, until = date.fromisoformat(out["period"]["since"]), date.fromisoformat(out["period"]["until"])
    by_id = {m["id"]: m for m in all_meetings(client)}
    assert {r["id"] for r in out["refs"]} <= set(by_id)
    for ref in out["refs"]:
        assert since <= date.fromisoformat(by_id[ref["id"]]["date"]) <= until
    assert sum(p["records"] for p in out["projects"]) == out["records"] == len(out["refs"])
    for p in out["projects"]:
        block = out["markdown"].split(f"## {p['title']}", 1)[1].split("\n## ", 1)[0]
        decisions = block.split("### Решения", 1)[1].split("###", 1)[0] if "### Решения" in block else ""
        assert len([l for l in decisions.splitlines() if re.match(r"\d+\. |- ", l)]) == min(p["decisions"], artifacts.MAX_SECTION)


def test_single_project_filter_and_errors(client):
    key, out = busy_period(client)
    project = next(p for p in out["projects"] if p["id"])
    one = digest_for(client, period=key, projectId=project["id"])
    assert [p["id"] for p in one["projects"]] == [project["id"]] and one["title"].endswith(project["title"])
    assert client.post("/api/materials/digest", json={"period": key, "projectId": "нет-такого"}).status_code == 404
    assert client.post("/api/materials/digest", json={"period": "year"}).status_code == 400


def test_empty_period_is_a_clear_error(client):
    core = client.app.state.core
    with pytest.raises(ValueError, match="записей нет"):
        projectdigest.collect(core, "this_week", services.meeting_date, today=date(2031, 3, 5))


def section(md: str, title: str) -> str:
    """Текст подраздела `### …` первого проекта дайджеста (пусто, если подраздела нет)."""
    if f"### {title}" not in md:
        return ""
    return md.split(f"### {title}", 1)[1].split("###", 1)[0].split("\n## ", 1)[0]


@pytest.fixture
def one_task(client):
    """Задача записи с проектом; её статус, срок и дату правки тест меняет и возвращает."""
    d = rich_meeting(client)
    task = d["tasks"][0]

    def set_state(**fields):
        with storage.connect() as db:
            for column, value in fields.items():
                db.execute(f"UPDATE tasks SET {column}=? WHERE id=?", (value, task["id"]))

    with storage.connect() as db:
        before = db.execute("SELECT status,deadline,updated_at FROM tasks WHERE id=?", (task["id"],)).fetchone()
    yield d, task, set_state
    set_state(status=before[0], deadline=before[1], updated_at=before[2])


def test_closed_in_period_section(client, one_task):
    d, task, set_state = one_task
    core, today = client.app.state.core, date.fromisoformat(d["date"])
    set_state(status="done", updated_at=today.isoformat() + "T10:00:00")
    md = projectdigest.build(core, "this_week", services.meeting_date, today=today)["markdown"]
    assert task["text"][:25] in section(md, "Закрыто за период")
    assert task["text"][:25] not in section(md, "Новые задачи")
    assert artifacts.untimed(md) == []
    # закрытая задолго до периода в «Закрыто за период» не попадает
    set_state(updated_at="2020-02-02T10:00:00")
    md = projectdigest.build(core, "this_week", services.meeting_date, today=today)["markdown"]
    assert task["text"][:25] not in section(md, "Закрыто за период")


def test_overdue_goes_to_attention_not_to_new_tasks(client, one_task):
    d, task, set_state = one_task
    core, today = client.app.state.core, date.fromisoformat(d["date"])
    set_state(status="open", deadline="2020-01-01")
    out = projectdigest.build(core, "this_week", services.meeting_date, today=today)
    md = out["markdown"]
    assert task["text"][:25] in section(md, "Требует внимания") and "просрочено" in section(md, "Требует внимания")
    assert task["text"][:25] not in section(md, "Новые задачи")
    assert sum(p["overdue"] for p in out["projects"]) >= 1
    assert artifacts.untimed(md) == []


def test_overdue_from_an_earlier_meeting_is_still_shown(client, one_task):
    """Просроченное из записи вне периода всё равно попадает в «Требует внимания» проекта, у которого в периоде есть запись."""
    d, task, set_state = one_task
    core = client.app.state.core
    set_state(status="open", deadline="2020-01-01")
    # «сегодня» сдвинут на неделю вперёд: запись задачи уже вне периода, а в периоде есть другая запись того же проекта
    project = d["project"]["id"] if isinstance(d["project"], dict) else None
    other = [m for m in core.live_meetings() if m.id != d["id"] and project and m.project_id == project]
    if not other:
        pytest.skip("в проекте записи только одна: сценарий проверяется на реальных данных")
    day = date.fromisoformat(services.meeting_date(other[0]))
    out = projectdigest.build(core, "this_week", services.meeting_date, today=day)
    if d["id"] in {r["id"] for r in out["refs"]}:
        pytest.skip("записи в одном периоде")
    assert task["text"][:25] in section(out["markdown"], "Требует внимания")


def test_quiet_projects_listed_without_bullets(client):
    core = client.app.state.core
    # сегодня далеко в будущем: записей в «этой неделе» нет → ошибка; а в «прошлой» неделе у фикстуры записи есть не всегда,
    # поэтому проверяем на каждом периоде: «Без записей за период» перечисляет только активные проекты вне дайджеста
    for key in KEYS:
        r = client.post("/api/materials/digest", json={"period": key})
        if r.status_code != 200:
            continue
        out = r.json()
        shown = {p["title"] for p in out["projects"]}
        if "## Без записей за период" in out["markdown"]:
            tail = out["markdown"].split("## Без записей за период", 1)[1].split("\n_", 1)[0]
            assert not any(t in tail for t in shown if t)
            assert not [l for l in tail.splitlines() if l.startswith(("- ", "|"))]
        # проекты из дайджеста не повторяются
        assert len(shown) == len(out["projects"])


# --------------------------------------------------------------------------- модель

def test_model_lead_and_fallbacks(client, monkeypatch):
    key, _ = busy_period(client)
    monkeypatch.setattr(digest, "generate", lambda *a, **k: "Команды продвинулись и согласовали порядок работ.")
    out = digest_for(client, period=key)
    assert out["mode"] == "model" and "Команды продвинулись и согласовали порядок работ." in out["markdown"]
    assert artifacts.untimed(out["markdown"]) == []
    monkeypatch.setattr(digest, "generate", lambda *a, **k: "Потратили 9 миллионов 4321 рубль.")
    bad = digest_for(client, period=key)
    assert bad["mode"] == "structured" and "4321" not in bad["markdown"]


# --------------------------------------------------------------------------- сохранение, экспорт, презентация, чат

def test_save_and_export(client):
    import os
    key, _ = busy_period(client)
    out = client.post("/api/materials/digest", json={"period": key, "save": True}).json()
    root = Path(os.environ["MEMORY_DESK_CONFIG"]).parent / "Claude" / "Дайджесты"
    assert (root / out["fileName"]).read_text(encoding="utf-8") == out["markdown"]
    for fmt in ("docx", "pdf", "html", "md"):
        r = client.post("/api/export", json={"format": fmt, "title": out["title"], "markdown": out["markdown"], "name": "Дайджест"})
        assert r.status_code == 200 and len(r.content) > 200
    assert client.post("/api/materials/save", json={"kind": "digest", "name": "Мой дайджест", "content": "# x\n"}).status_code == 200


def test_deck(client):
    from pptx import Presentation
    key, out = busy_period(client)
    r = client.post("/api/materials/digest/deck", json={"period": key})
    assert r.status_code == 200, r.text
    info = r.json()
    assert info["slides"] >= 3 and info["url"].endswith(".pptx")
    prs = Presentation(BytesIO(client.get(info["url"]).content))
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Дайджест" in texts and out["period"]["label"] in texts
    for p in out["projects"]:
        assert p["title"] in texts or p["title"][:28] in texts
    assert client.post("/api/materials/digest/deck", json={"period": "year"}).status_code == 400


def test_chat_digest(client):
    ask = lambda q: client.post("/api/chat", json={"question": q, "scope": {"kind": "all"}}).json()
    reply = None
    for phrase in ("сделай дайджест за эту неделю", "сделай дайджест за прошлую неделю", "сделай месячный дайджест за этот месяц", "сделай месячный дайджест за прошлый месяц"):
        reply = ask(phrase)
        if "material" in reply:
            break
    assert reply and reply["material"]["kind"] == "digest" and reply["material"]["result"]["markdown"].startswith("# Дайджест")
    deck = None
    for phrase in ("подготовь презентацию-дайджест за этот месяц", "подготовь презентацию-дайджест за прошлый месяц", "подготовь презентацию-дайджест за эту неделю", "подготовь презентацию-дайджест за прошлую неделю"):
        deck = ask(phrase)
        if "material" in deck:
            break
    assert deck["material"]["kind"] == "deck" and deck["material"]["result"]["url"].endswith(".pptx")
    # просьба без глагола — обычный вопрос
    assert "material" not in ask("какой дайджест мне нужен?")


def test_request_kind_digest():
    k = materials.request_kind
    assert k("подготовь недельный дайджест") == "digest" and k("сделай дайджест в виде презентации") == "digest"
    assert k("сделай презентацию") == "deck" and k("составь протокол") == "protocol"
